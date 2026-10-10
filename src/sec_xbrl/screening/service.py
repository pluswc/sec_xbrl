"""Deterministic screening over immutable prepared screening datasets."""

from __future__ import annotations

import csv
import hashlib
import json
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sec_xbrl.screening.common import (
    canonical_json,
    digest_value,
    require_cik,
    require_string,
    safe_new_destination,
    sha256_file,
)
from sec_xbrl.screening.company_descriptions import CompanyDescriptionReader
from sec_xbrl.screening.dataset import ScreeningDatasetReader
from sec_xbrl.screening.dataset_v2 import ScreeningDatasetV2Reader, open_screening_dataset
from sec_xbrl.screening.dataset_v3 import ScreeningDatasetV3Reader
from sec_xbrl.screening.explorer import ScreeningExplorer, public_copy
from sec_xbrl.screening.store import ScreeningStore

ENGINE_VERSION = "screening-executor-v1"
EXPORT_VERSION = "screen-run-export-v1"
XLSX_LONG_TEXT_VERSION = "screen-run-xlsx-long-text-v1"
XLSX_DIRECT_TEXT_LIMIT = 32_000
XLSX_CHUNK_TEXT_LIMIT = 16_000
OPERATORS = {"GT", "GTE", "LT", "LTE", "EQ"}
AGGREGATES = {"PASS", "FAIL", "UNDETERMINED"}
VISIBLE_COLUMNS = {
    "company",
    "ticker",
    "name",
    "cik",
    "sic",
    "sic_code",
    "aggregate",
    "predicates",
    "decision",
}


class ScreeningService:
    """Application service over verified publications and an append-only store."""

    def __init__(
        self,
        datasets: dict[str, Path | str],
        store: ScreeningStore | Path | str,
        *,
        company_descriptions: CompanyDescriptionReader | None = None,
    ) -> None:
        if not isinstance(datasets, dict) or not datasets:
            raise ValueError("at least one screening dataset is required")
        self.datasets: dict[str, ScreeningDatasetReader | ScreeningDatasetV2Reader] = {}
        self.explorers: dict[str, ScreeningExplorer] = {}
        for dataset_id, root in datasets.items():
            dataset_id = require_string(dataset_id, "dataset catalog ID")
            reader = open_screening_dataset(root)
            if reader.manifest["dataset_id"] != dataset_id:
                raise ValueError("dataset catalog ID differs from publication")
            self.datasets[dataset_id] = reader
            if isinstance(reader, (ScreeningDatasetV2Reader, ScreeningDatasetV3Reader)):
                self.explorers[dataset_id] = ScreeningExplorer(reader)
        if company_descriptions is not None:
            bound = self.datasets.get(company_descriptions.dataset_id)
            if bound is None or bound.publication_id != company_descriptions.dataset.publication_id:
                raise ValueError("company descriptions differ from the screening catalog")
        self.company_descriptions = company_descriptions
        self.store = store if isinstance(store, ScreeningStore) else ScreeningStore(store)

    def status(self) -> dict[str, Any]:
        status = {
            "engine_version": ENGINE_VERSION,
            "datasets": [reader.verify_all() for reader in self.datasets.values()],
            "screens": len(self.store.list_screens()),
            "runs": len(self.store.list_runs()),
        }
        if self.company_descriptions is not None:
            status["company_descriptions"] = self.company_descriptions.verify_all()
        return status

    def metric_availability(self, dataset_id: str, *, cik: str | None = None) -> list[dict[str, Any]]:
        reader = self._dataset(dataset_id)
        rows = reader.metric_support
        if cik is not None:
            cik = require_cik(cik)
            rows = [row for row in rows if row["cik"] == cik]
        return json.loads(canonical_json(rows))

    def metric_catalog(self, dataset_id: str) -> list[dict[str, Any]]:
        reader = self._dataset(dataset_id)
        definitions = getattr(reader, "metric_definitions", None)
        if definitions is None:
            labels: dict[str, str] = {}
            for row in reader.metric_support:
                labels.setdefault(row["metric_id"], row.get("label") or row["metric_id"])
            definitions = [
                {
                    "metric_id": metric_id,
                    "label": label,
                    "category": "LEGACY_REPORTED",
                    "unit_kind": "MONEY",
                    "unit": "USD",
                    "definition_version": next(
                        row["definition_version"]
                        for row in reader.metric_support
                        if row["metric_id"] == metric_id
                    ),
                    "period_class": "FY",
                    "source_kind": "REPORTED",
                }
                for metric_id, label in sorted(labels.items())
            ]
        result = []
        for definition in definitions:
            rows = [row for row in reader.metric_support if row["metric_id"] == definition["metric_id"]]
            result.append(
                {
                    **json.loads(canonical_json(definition)),
                    "coverage": {
                        "companies_available": sum(row.get("execution_supported") is True for row in rows),
                        "companies_total": len(rows),
                        "periods_available": sum(int(row.get("available_periods") or 0) for row in rows),
                    },
                    "business_availability": "AVAILABLE" if any(row.get("execution_supported") for row in rows) else "UNAVAILABLE",
                    "unavailable_reasons": sorted(
                        {
                            reason
                            for row in rows
                            for reason in (row.get("reasons") or [row.get("reason")])
                            if reason
                        }
                    ),
                }
            )
        return result

    def quarter_metric_catalog(self, dataset_id: str) -> list[dict[str, Any]]:
        reader = self._quarter_dataset(dataset_id)
        result = []
        for definition in reader.quarter_definitions:
            rows = [row for row in reader.quarter_observations if row["metric_id"] == definition["metric_id"]]
            result.append({**json.loads(canonical_json(definition)), "label": definition["label_ko"], "period_class": definition["observation_basis"], "coverage": {"companies_available": len({row["cik"] for row in rows if row["availability_status"] == "AVAILABLE"}), "companies_total": len(reader.companies), "periods_available": sum(row["availability_status"] == "AVAILABLE" for row in rows)}})
        return result

    def quarter_periods(self, dataset_id: str, *, cik: str | None = None) -> dict[str, Any]:
        reader = self._quarter_dataset(dataset_id)
        if cik is not None:
            cik = require_cik(cik)
        rows = [row for row in reader.quarter_period_inventory if cik is None or row["cik"] == cik]
        return {"spine": json.loads(canonical_json(reader.quarterly.manifest["global_quarter_spine"])), "rows": json.loads(canonical_json(rows))}

    def quarter_series(self, dataset_id: str, *, cik: str, metric_id: str, limit: int = 8) -> dict[str, Any]:
        reader = self._quarter_dataset(dataset_id); cik = require_cik(cik); metric_id = require_string(metric_id, "metric_id")
        if type(limit) is not int or not 1 <= limit <= 8: raise ValueError("quarter series limit must be one to eight")
        rows = [row for row in reader.quarter_observations if row["cik"] == cik and row["metric_id"] == metric_id]
        if not rows: raise ValueError("quarter series is not prepared")
        rows.sort(key=lambda row: (row["calendar_year"], row["calendar_quarter"]))
        return {"metric": next(row for row in reader.quarter_definitions if row["metric_id"] == metric_id), "observations": json.loads(canonical_json(rows[-limit:]))}

    def quarter_comparison(self, dataset_id: str, *, ciks: list[str], metric_id: str, calendar_year: int, calendar_quarter: int) -> dict[str, Any]:
        reader = self._quarter_dataset(dataset_id)
        if len(ciks) < 2 or len(ciks) > 4 or len(set(ciks)) != len(ciks): raise ValueError("quarter comparison requires two to four unique companies")
        rows=[]
        for raw in ciks:
            cik=require_cik(raw); company=reader.company(cik); matches=[row for row in reader.quarter_observations if row["cik"]==cik and row["metric_id"]==metric_id and row["calendar_year"]==calendar_year and row["calendar_quarter"]==calendar_quarter]
            if len(matches)!=1: raise ValueError("quarter comparison cell is not unique")
            rows.append({"company":company,"observation":json.loads(canonical_json(matches[0]))})
        signatures={(r["observation"]["definition_version"],r["observation"]["unit"],r["observation"]["selection_as_of_date"]) for r in rows}
        if len(signatures)!=1: raise ValueError("quarter comparison scope is incompatible")
        return {"calendar_year":calendar_year,"calendar_quarter":calendar_quarter,"metric_id":metric_id,"rows":rows}

    def company_explorer(
        self,
        dataset_id: str,
        *,
        fiscal_year: int | None = None,
        query: str | None = None,
        sort: str = "company",
        direction: str = "asc",
    ) -> dict[str, Any]:
        """Return the company browse matrix without creating a screen or run."""
        return public_copy(
            self._explorer(dataset_id).explore(
                fiscal_year=fiscal_year,
                query=query,
                sort=sort,
                direction=direction,
            )
        )

    def metric_library(
        self,
        dataset_id: str,
        *,
        fiscal_year: int | None = None,
        query: str | None = None,
        category: str | None = None,
    ) -> dict[str, Any]:
        """Return selected-year coverage and guidance for every governed metric."""
        return public_copy(
            self._explorer(dataset_id).metric_library(
                fiscal_year=fiscal_year,
                query=query,
                category=category,
            )
        )

    def company_profile(
        self, dataset_id: str, *, cik: str, fiscal_year: int | None = None
    ) -> dict[str, Any]:
        """Return a run-independent company profile over prepared observations."""
        result = public_copy(
            self._explorer(dataset_id).company_profile(
                cik=cik,
                fiscal_year=fiscal_year,
            )
        )
        if self.company_descriptions is not None and self.company_descriptions.dataset_id == dataset_id:
            row = self.company_descriptions.company(cik)
            result["company_description"] = {
                key: row.get(key)
                for key in (
                    "status", "text", "locale", "provider_name", "profile_url",
                    "checked_at", "observed_at", "publication_id", "temporal_scope", "reason",
                    "content_fixture_kind",
                )
            }
            result["company_description"]["publication_id"] = self.company_descriptions.publication_id
        return result

    def company_comparison(
        self,
        dataset_id: str,
        *,
        ciks: list[str],
        metric_ids: list[str],
        fiscal_year: int | None = None,
    ) -> dict[str, Any]:
        """Return a same-year matrix for two to four prepared companies."""
        return public_copy(
            self._explorer(dataset_id).compare(
                ciks=ciks,
                metric_ids=metric_ids,
                fiscal_year=fiscal_year,
            )
        )

    def metric_series(self, dataset_id: str, *, cik: str, metric_id: str) -> dict[str, Any]:
        reader = self._dataset(dataset_id)
        cik = require_cik(cik)
        metric_id = require_string(metric_id, "metric_id")
        if not any(row["metric_id"] == metric_id for row in reader.metric_support):
            raise ValueError("unknown metric_id")
        rows = [
            row for row in reader.observations
            if row["cik"] == cik and row["metric_id"] == metric_id
        ]
        if not rows:
            raise ValueError("metric is not prepared for this company")
        rows.sort(key=lambda row: (row["fiscal_year"], row["period_class"], row["view"]))
        return {"company": reader.company(cik), "metric_id": metric_id, "observations": rows}

    def metadata_inventory(
        self,
        dataset_id: str,
        *,
        query: str | None = None,
        sic_code: str | None = None,
        prepared: bool | None = None,
        offset: int = 0,
        limit: int = 100,
    ) -> dict[str, Any]:
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 500:
            raise ValueError("invalid metadata page")
        rows = self._dataset(dataset_id).metadata
        if query:
            needle = query.casefold()
            rows = [
                row for row in rows
                if needle in str(row.get("name") or "").casefold()
                or any(needle in ticker.casefold() for ticker in row.get("tickers", []))
                or needle in row["cik"]
            ]
        if sic_code is not None:
            rows = [row for row in rows if row.get("sic_code") == sic_code]
        if prepared is not None:
            rows = [row for row in rows if row.get("financial_prepared") is prepared]
        return {"total": len(rows), "offset": offset, "limit": limit, "rows": rows[offset:offset + limit]}

    def peer_candidates(self, dataset_id: str, *, cik: str) -> dict[str, Any]:
        reader = self._dataset(dataset_id)
        cik = require_cik(cik)
        target = next((row for row in reader.metadata if row["cik"] == cik), None)
        if target is None:
            raise ValueError("company is not present in metadata inventory")
        sic = target.get("sic_code")
        peers = [] if sic is None else [
            row for row in reader.metadata if row["cik"] != cik and row.get("sic_code") == sic
        ]
        return {
            "target": target,
            "peers": peers,
            "limitation": "SEC SIC is a current coarse peer-candidate signal, not accounting equivalence.",
        }

    def create_screen(self, name: str, content: dict[str, Any]) -> dict[str, Any]:
        self._validate_revision(content)
        return self.store.create_screen(name, content)

    def revise_screen(self, screen_id: str, content: dict[str, Any]) -> dict[str, Any]:
        self._validate_revision(content)
        return self.store.revise_screen(screen_id, content)

    def execute(self, revision_id: str, dataset_id: str) -> dict[str, Any]:
        reader = self._dataset(dataset_id)
        revision = self.store.revision(revision_id)
        content = revision["content"]
        self._validate_revision(content, reader=reader)
        targets = _select_universe(reader, content["universe"])
        universe_payload = {
            "dataset_publication_id": reader.publication_id,
            "metadata_publication_id": reader.manifest["sources"]["metadata"]["publication_id"],
            "selector": content["universe"],
            "ciks": [row["cik"] for row in targets],
        }
        universe = self.store.save_universe(universe_payload)
        rows = [self._evaluate_company(reader, company, content) for company in targets]
        rows = _sort_rows(rows, content["sort"])
        counts = {state.lower(): sum(row["aggregate"] == state for row in rows) for state in AGGREGATES}
        counts.update(
            target=len(rows),
            quality_block=sum(row["quality_blocking"] for row in rows),
        )
        if counts["target"] != counts["pass"] + counts["fail"] + counts["undetermined"]:
            raise AssertionError("screen result counts do not reconcile")
        dataset_versions = json.loads(canonical_json(reader.manifest.get("versions", {})))
        predicate_versions = sorted(
            {
                predicate["definition_version"]
                for predicate in content["predicates"]
            }
        )
        versions = {
            "engine": ENGINE_VERSION,
            "dataset_publication_id": reader.publication_id,
            "dataset": dataset_versions,
            "predicate_definition_versions": predicate_versions,
            "metadata_publication_id": reader.manifest["sources"]["metadata"]["publication_id"],
        }
        input_material = {
            "revision_id": revision_id,
            "revision_hash": revision["content_hash"],
            "universe_snapshot_id": universe["universe_snapshot_id"],
            "dataset_publication_id": reader.publication_id,
            "versions": versions,
        }
        payload = {
            "revision_id": revision_id,
            "universe_snapshot_id": universe["universe_snapshot_id"],
            "dataset_id": dataset_id,
            "dataset_publication_id": reader.publication_id,
            "fixture_kind": reader.manifest["fixture_kind"],
            "financial_as_of": reader.manifest["financial_as_of"],
            "period_policy": content["period_policy"],
            "versions": versions,
            "input_hash": digest_value(input_material),
            "counts": counts,
            "rows": rows,
        }
        return self.store.save_run(payload)

    def run(self, run_id: str) -> dict[str, Any]:
        return self.store.run(run_id)

    def trace(self, run_id: str, *, cik: str, metric_id: str) -> dict[str, Any]:
        run = self.store.run(run_id)
        cik = require_cik(cik)
        metric_id = require_string(metric_id, "metric_id")
        row = next((row for row in run["rows"] if row["cik"] == cik), None)
        if row is None:
            raise ValueError("company is not in the run")
        evaluation = next(
            (item for item in row["predicates"] if item["metric_id"] == metric_id), None
        )
        if evaluation is None:
            raise ValueError("metric is not in the run")
        return {
            "run_id": run_id,
            "company": {key: row[key] for key in ("cik", "ticker", "name", "sic_code")},
            "evaluation": evaluation,
            "trace": evaluation.get("observation", {}).get("trace"),
        }

    def compare_peers(
        self,
        dataset_id: str,
        *,
        ciks: list[str],
        metric_id: str,
        fiscal_year: int,
    ) -> dict[str, Any]:
        if len(ciks) < 2 or len(ciks) > 100 or len(ciks) != len(set(ciks)):
            raise ValueError("peer comparison requires two to 100 unique CIKs")
        if type(fiscal_year) is not int or not 1900 <= fiscal_year <= 2200:
            raise ValueError("peer comparison fiscal year is invalid")
        reader = self._dataset(dataset_id)
        definition_versions = {
            row.get("definition_version")
            for row in reader.metric_support
            if row.get("metric_id") == metric_id and row.get("execution_supported") is True
        }
        if len(definition_versions) != 1:
            raise ValueError("peer comparison metric is unavailable or version-ambiguous")
        definition_version = next(iter(definition_versions))
        rows = []
        for cik in ciks:
            cik = require_cik(cik)
            company = reader.company(cik)
            observation = _find_observation(
                reader,
                cik,
                metric_id,
                {"mode": "FIXED_FISCAL_YEAR", "fiscal_year": fiscal_year},
                definition_version=definition_version,
            )
            rows.append({"company": company, "observation": observation})
        return {
            "metric_id": metric_id,
            "fiscal_year": fiscal_year,
            "rows": rows,
            "warning": "같은 FY label은 실제 시작일·context end 날짜의 동등성을 뜻하지 않습니다.",
        }

    def add_decision(self, **kwargs: Any) -> dict[str, Any]:
        return self.store.add_decision(**kwargs)

    def decisions(self, cik: str | None = None) -> list[dict[str, Any]]:
        return self.store.decisions(cik)

    def compare_runs(self, old_run_id: str, new_run_id: str) -> dict[str, Any]:
        old, new = self.store.run(old_run_id), self.store.run(new_run_id)
        old_universe = self.store.universe(old["universe_snapshot_id"])["content"]
        new_universe = self.store.universe(new["universe_snapshot_id"])["content"]
        old_rule_vector = {
            "engine": old["versions"].get("engine"),
            "dataset": old["versions"].get("dataset"),
            "predicate_definition_versions": old["versions"].get(
                "predicate_definition_versions"
            ),
        }
        new_rule_vector = {
            "engine": new["versions"].get("engine"),
            "dataset": new["versions"].get("dataset"),
            "predicate_definition_versions": new["versions"].get(
                "predicate_definition_versions"
            ),
        }
        changed = {
            "data": old["dataset_publication_id"] != new["dataset_publication_id"],
            "rule": old_rule_vector != new_rule_vector,
            "universe": old_universe.get("ciks") != new_universe.get("ciks"),
            "universe_selector": old_universe.get("selector") != new_universe.get("selector"),
            "condition_revision": old["revision_id"] != new["revision_id"],
            "manual_decision": _decision_state(self.store.decisions(), old_run_id)
            != _decision_state(self.store.decisions(), new_run_id),
        }
        changed_dimensions = [name for name, value in changed.items() if value]
        attribution = (
            "NO_VERSION_CHANGE"
            if not changed_dimensions
            else changed_dimensions[0].upper() + "_CHANGE"
            if len(changed_dimensions) == 1
            else "MIXED_OR_INCOMPARABLE_VERSION_CHANGE"
        )
        old_rows = {row["cik"]: row for row in old["rows"]}
        new_rows = {row["cik"]: row for row in new["rows"]}
        differences = [
            {
                "cik": cik,
                "old": old_rows.get(cik),
                "new": new_rows.get(cik),
            }
            for cik in sorted(set(old_rows) | set(new_rows))
            if old_rows.get(cik) != new_rows.get(cik)
        ]
        return {
            "old_run_id": old_run_id,
            "new_run_id": new_run_id,
            "changed_dimensions": changed,
            "attribution": attribution,
            "differences": differences,
            "result_rows_changed": bool(differences),
            "result_state": "CHANGED" if differences else "UNCHANGED",
            "fixture_kinds": [old["fixture_kind"], new["fixture_kind"]],
            "universe_membership": {
                "old_ciks": old_universe.get("ciks", []),
                "new_ciks": new_universe.get("ciks", []),
            },
            "version_vectors": {"old": old_rule_vector, "new": new_rule_vector},
        }

    def export_run(
        self,
        run_id: str,
        destination: Path | str,
        *,
        scope: str = "FULL",
        selected_ciks: list[str] | None = None,
        format: str | None = None,
    ) -> dict[str, Any]:
        run = self.store.run(run_id)
        if scope not in {"FULL", "SELECTED"}:
            raise ValueError("export scope must be FULL or SELECTED")
        selected = None
        if scope == "SELECTED":
            if not selected_ciks:
                raise ValueError("selected export requires CIKs")
            selected = [require_cik(cik) for cik in selected_ciks]
            if len(selected) != len(set(selected)):
                raise ValueError("selected export contains duplicate CIKs")
        rows = run["rows"] if selected is None else [row for row in run["rows"] if row["cik"] in selected]
        if selected is not None and {row["cik"] for row in rows} != set(selected):
            raise ValueError("selected export contains a CIK outside the run")
        destination_path = safe_new_destination(Path(destination), "screening export")
        selected_format = (format or destination_path.suffix.lstrip(".")).lower()
        if selected_format not in {"json", "csv", "xlsx"}:
            raise ValueError("screening export format must be JSON, CSV, or XLSX")
        payload = {
            "export_version": EXPORT_VERSION,
            "scope": scope,
            "selected_ciks": selected,
            "run": {key: value for key, value in run.items() if key != "rows"},
            "rows": rows,
        }
        destination_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination_path.parent / (
            f".{destination_path.name}.partial-{uuid.uuid4().hex}"
        )
        try:
            if selected_format == "json":
                temporary.write_text(canonical_json(payload) + "\n", encoding="utf-8")
            elif selected_format == "csv":
                _write_csv(temporary, payload)
            else:
                _write_xlsx(temporary, payload)
            temporary.rename(destination_path)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        manifest = {
            "export_version": EXPORT_VERSION,
            "run_id": run_id,
            "screen_revision_id": run["revision_id"],
            "universe_snapshot_id": run["universe_snapshot_id"],
            "dataset_publication_id": run["dataset_publication_id"],
            "financial_as_of": run["financial_as_of"],
            "versions": run["versions"],
            "scope": scope,
            "selected_ciks": selected,
            "row_count": len(rows),
            "format": selected_format,
            "path": str(destination_path),
            "sha256": sha256_file(destination_path),
            "size": destination_path.stat().st_size,
            "canonical_payload_hash": digest_value(payload),
        }
        if selected_format == "xlsx":
            manifest["xlsx_encoding_version"] = XLSX_LONG_TEXT_VERSION
        try:
            return self.store.record_export(manifest)
        except Exception:
            destination_path.unlink(missing_ok=True)
            raise

    def _dataset(self, dataset_id: str) -> ScreeningDatasetReader | ScreeningDatasetV2Reader:
        try:
            return self.datasets[require_string(dataset_id, "dataset_id")]
        except KeyError as exc:
            raise ValueError("unknown screening dataset") from exc

    def _explorer(self, dataset_id: str) -> ScreeningExplorer:
        dataset_id = require_string(dataset_id, "dataset_id")
        if dataset_id not in self.datasets:
            raise ValueError("unknown screening dataset")
        try:
            return self.explorers[dataset_id]
        except KeyError as exc:
            raise ValueError("company explorer requires a v2 screening dataset") from exc

    def _quarter_dataset(self, dataset_id: str) -> ScreeningDatasetV3Reader:
        reader = self._dataset(dataset_id)
        if not isinstance(reader, ScreeningDatasetV3Reader):
            raise TypeError("quarterly analysis requires a v3 screening dataset")
        return reader

    def _validate_revision(
        self, content: dict[str, Any], *, reader: ScreeningDatasetReader | ScreeningDatasetV2Reader | None = None
    ) -> None:
        if not isinstance(content, dict):
            raise TypeError("screen revision must be an object")
        predicates = content.get("predicates")
        if not isinstance(predicates, list) or not predicates:
            raise ValueError("screen revision requires predicates")
        ids = []
        for predicate in predicates:
            if not isinstance(predicate, dict):
                raise TypeError("predicate must be an object")
            metric_id = require_string(predicate.get("metric_id"), "metric_id")
            ids.append(metric_id)
            if predicate.get("operator") not in OPERATORS:
                raise ValueError("unsupported predicate operator")
            _decimal(predicate.get("threshold"))
            definition_version = require_string(
                predicate.get("definition_version"), "metric definition version"
            )
            if reader is not None:
                quarterly = content.get("period_policy", {}).get("mode") == "FIXED_CALENDAR_QUARTER_END"
                supports = [row for row in (getattr(reader, "quarter_support", []) if quarterly else reader.metric_support) if row["metric_id"] == metric_id]
                if not supports:
                    raise ValueError("unknown metric_id")
                if not any(row.get("execution_supported", row.get("available_bucket_count", 0) > 0) for row in supports):
                    reason = sorted({row.get("reason") for row in supports})
                    raise ValueError(f"metric is unavailable for execution: {reason}")
                if definition_version not in {
                    row.get("definition_version") for row in supports
                }:
                    raise ValueError("metric definition version is not prepared")
        if len(ids) != len(set(ids)):
            raise ValueError("duplicate metric predicate")
        _validate_universe(content.get("universe"))
        _validate_period_policy(content.get("period_policy"))
        columns = content.get("columns")
        if (
            not isinstance(columns, list)
            or not columns
            or any(value not in VISIBLE_COLUMNS for value in columns)
            or len(columns) != len(set(columns))
        ):
            raise ValueError("screen columns must be a nonempty string list")
        sort = content.get("sort")
        if not isinstance(sort, dict) or sort.get("direction") not in {"ASC", "DESC"}:
            raise ValueError("invalid screen sort")
        sort_column = require_string(sort.get("column"), "sort column")
        if sort_column.startswith("metric:") and sort_column.split(":", 1)[1] not in ids:
            raise ValueError("metric sort column requires a matching predicate")

    def _evaluate_company(
        self, reader: ScreeningDatasetReader | ScreeningDatasetV2Reader, company: dict[str, Any], content: dict[str, Any]
    ) -> dict[str, Any]:
        evaluations = []
        for predicate in content["predicates"]:
            observation = _find_observation(
                reader,
                company["cik"],
                predicate["metric_id"],
                content["period_policy"],
                definition_version=predicate["definition_version"],
            )
            evaluations.append(_evaluate_predicate(predicate, observation))
        valid_fail = any(row["result"] == "FAIL" for row in evaluations)
        quality_blocking = any(row["quality_blocking"] for row in evaluations)
        if valid_fail:
            aggregate = "FAIL"
        elif all(row["result"] == "PASS" for row in evaluations) and not quality_blocking:
            aggregate = "PASS"
        else:
            aggregate = "UNDETERMINED"
        return {
            "cik": company["cik"],
            "ticker": company["ticker"],
            "name": company["name"],
            "sic_code": company.get("sic_code"),
            "sic_description": company.get("sic_description"),
            "aggregate": aggregate,
            "quality_blocking": quality_blocking,
            "predicates": evaluations,
        }


def _select_universe(reader: ScreeningDatasetReader | ScreeningDatasetV2Reader, selector: dict[str, Any]) -> list[dict[str, Any]]:
    mode = selector["mode"]
    rows = reader.companies
    if mode == "CIKS":
        selected = set(selector["ciks"])
        rows = [row for row in rows if row["cik"] in selected]
        if {row["cik"] for row in rows} != selected:
            raise ValueError("universe CIK is not financially prepared")
    elif mode == "SIC":
        rows = [row for row in rows if row.get("sic_code") in set(selector["sic_codes"])]
    if not rows:
        raise ValueError("universe selector has no prepared target")
    return sorted(rows, key=lambda row: row["cik"])


def _find_observation(
    reader: ScreeningDatasetReader | ScreeningDatasetV2Reader,
    cik: str,
    metric_id: str,
    policy: dict[str, Any],
    *,
    definition_version: str | None = None,
) -> dict[str, Any] | None:
    quarterly = policy["mode"] == "FIXED_CALENDAR_QUARTER_END"
    source_rows = getattr(reader, "quarter_observations", []) if quarterly else reader.observations
    rows = [row for row in source_rows if row["cik"] == cik and row["metric_id"] == metric_id]
    if definition_version is not None:
        rows = [
            row
            for row in rows
            if row.get("metric_definition_version", row.get("definition_version")) == definition_version
        ]
    if policy["mode"] == "FIXED_FISCAL_YEAR":
        rows = [row for row in rows if row["fiscal_year"] == policy["fiscal_year"]]
    elif quarterly:
        rows = [row for row in rows if row["calendar_year"] == policy["calendar_year"] and row["calendar_quarter"] == policy["calendar_quarter"]]
        rows = [{**row, "quality_blocking": False, "trace": row.get("source_lineage")} for row in rows]
    else:
        rows.sort(key=lambda row: row["fiscal_year"], reverse=True)
        offset = policy.get("offset", 0)
        rows = rows[offset:offset + 1]
    if len(rows) > 1:
        raise ValueError("metric observation is not unique at screening grain")
    return json.loads(canonical_json(rows[0])) if rows else None


def _evaluate_predicate(
    predicate: dict[str, Any], observation: dict[str, Any] | None
) -> dict[str, Any]:
    base = {
        "metric_id": predicate["metric_id"],
        "definition_version": predicate["definition_version"],
        "operator": predicate["operator"],
        "threshold": str(_decimal(predicate["threshold"])),
        "observation": observation,
        "quality_blocking": bool(observation and observation.get("quality_blocking")),
    }
    if observation is None:
        return {**base, "result": "UNDETERMINED", "reason": "PERIOD_NOT_PREPARED"}
    if observation["availability_status"] != "AVAILABLE" or observation.get("quality_blocking"):
        return {
            **base,
            "result": "UNDETERMINED",
            "reason": observation.get("unavailable_reason") or "OBSERVATION_UNAVAILABLE",
        }
    value = _decimal(observation["value"])
    threshold = _decimal(predicate["threshold"])
    operator = predicate["operator"]
    passed = {
        "GT": value > threshold,
        "GTE": value >= threshold,
        "LT": value < threshold,
        "LTE": value <= threshold,
        "EQ": value == threshold,
    }[operator]
    return {
        **base,
        "result": "PASS" if passed else "FAIL",
        "reason": "PREDICATE_MATCHED" if passed else "PREDICATE_NOT_MATCHED",
    }


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError("numeric value must be a finite decimal")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("numeric value must be a finite decimal") from exc
    if not result.is_finite():
        raise ValueError("numeric value must be a finite decimal")
    return result


def _validate_universe(value: Any) -> None:
    if not isinstance(value, dict) or value.get("mode") not in {"ALL_PREPARED", "CIKS", "SIC"}:
        raise ValueError("invalid universe selector")
    mode = value["mode"]
    allowed = {"mode"} | ({"ciks"} if mode == "CIKS" else {"sic_codes"} if mode == "SIC" else set())
    if set(value) != allowed:
        raise ValueError("universe selector contains unexpected fields")
    if mode == "CIKS":
        rows = value.get("ciks")
        if not isinstance(rows, list) or not rows:
            raise ValueError("CIK universe requires CIKs")
        ciks = [require_cik(cik) for cik in rows]
        if len(ciks) != len(set(ciks)):
            raise ValueError("duplicate universe CIK")
    if mode == "SIC":
        rows = value.get("sic_codes")
        if not isinstance(rows, list) or not rows or any(not isinstance(row, str) for row in rows):
            raise ValueError("SIC universe requires string SIC codes")
        if len(rows) != len(set(rows)):
            raise ValueError("duplicate universe SIC")


def _validate_period_policy(value: Any) -> None:
    modes = {"FIXED_FISCAL_YEAR", "LATEST_COMPLETE_ANNUAL", "FIXED_CALENDAR_QUARTER_END"}
    if not isinstance(value, dict) or value.get("mode") not in modes:
        raise ValueError("invalid period policy")
    if value["mode"] == "FIXED_FISCAL_YEAR":
        year = value.get("fiscal_year")
        if type(year) is not int or not 1900 <= year <= 2200 or set(value) != {"mode", "fiscal_year"}:
            raise ValueError("fixed annual policy requires one fiscal year")
    elif value["mode"] == "LATEST_COMPLETE_ANNUAL":
        offset = value.get("offset", 0)
        if type(offset) is not int or not 0 <= offset <= 20 or set(value) - {"mode", "offset"}:
            raise ValueError("latest annual offset is invalid")
    else:
        year, quarter = value.get("calendar_year"), value.get("calendar_quarter")
        if (
            type(year) is not int
            or not 1900 <= year <= 2200
            or type(quarter) is not int
            or quarter not in {1, 2, 3, 4}
            or set(value) != {"mode", "calendar_year", "calendar_quarter"}
        ):
            raise ValueError("fixed calendar quarter-end policy requires one year and quarter")


def _sort_rows(rows: list[dict[str, Any]], spec: dict[str, Any]) -> list[dict[str, Any]]:
    column = require_string(spec.get("column"), "sort column")
    direction = spec["direction"]
    if column not in {"cik", "ticker", "name", "aggregate"} and not column.startswith("metric:"):
        raise ValueError("unsupported sort column")

    def value(row: dict[str, Any]) -> tuple[bool, Any]:
        if column.startswith("metric:"):
            metric_id = column.split(":", 1)[1]
            item = next((item for item in row["predicates"] if item["metric_id"] == metric_id), None)
            raw = None if item is None or item["observation"] is None else item["observation"].get("value")
            return raw is None, Decimal(raw) if raw is not None else Decimal(0)
        raw = row.get(column)
        return raw is None, str(raw or "")

    by_cik = sorted(rows, key=lambda row: row["cik"])
    present = [row for row in by_cik if not value(row)[0]]
    missing = [row for row in by_cik if value(row)[0]]
    present.sort(key=lambda row: value(row)[1], reverse=direction == "DESC")
    return [*present, *missing]


def _decision_state(events: list[dict[str, Any]], run_id: str) -> dict[str, str]:
    result = {}
    for event in events:
        if event["run_id"] == run_id:
            result[event["cik"]] = event["state"]
    return result


def _write_csv(path: Path, payload: dict[str, Any]) -> None:
    fields = ["record_type", "cik", "ticker", "name", "sic_code", "aggregate", "quality_blocking", "record_json"]
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerow(
            {
                "record_type": "EXPORT",
                "record_json": _csv_safe(
                    {
                        "metadata": {
                            key: value for key, value in payload.items() if key != "rows"
                        },
                        "canonical_payload_hash": digest_value(payload),
                    }
                ),
            }
        )
        for row in payload["rows"]:
            writer.writerow({
                "record_type": "ROW",
                "cik": row["cik"],
                "ticker": _csv_safe(row["ticker"]),
                "name": _csv_safe(row.get("name")),
                "sic_code": row.get("sic_code"),
                "aggregate": row["aggregate"],
                "quality_blocking": str(row["quality_blocking"]).lower(),
                "record_json": _csv_safe(row),
            })


def _write_xlsx(path: Path, payload: dict[str, Any]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    rows_sheet = workbook.active
    rows_sheet.title = "Rows"
    long_records: list[dict[str, Any]] = []
    headers = ["cik", "ticker", "name", "sic_code", "aggregate", "quality_blocking", "record_json"]
    rows_sheet.append(headers)
    for index, row in enumerate(payload["rows"]):
        rows_sheet.append(
            [
                str(row.get(key, ""))
                if key != "record_json"
                else _xlsx_text(
                    canonical_json(row),
                    record_key=f"row:{index}:{row['cik']}",
                    long_records=long_records,
                )
                for key in headers
            ]
        )
    metadata = workbook.create_sheet("Metadata")
    metadata.append(["key", "value"])
    for key, value in (
        ("export_version", payload["export_version"]),
        ("xlsx_long_text_encoding", XLSX_LONG_TEXT_VERSION),
        ("scope", payload["scope"]),
        ("selected_ciks_json", canonical_json(payload["selected_ciks"])),
        ("run_json", canonical_json(payload["run"])),
        ("canonical_payload_hash", digest_value(payload)),
    ):
        metadata.append(
            [
                key,
                _xlsx_text(
                    value,
                    record_key=f"metadata:{key}",
                    long_records=long_records,
                ),
            ]
        )
    long_text = workbook.create_sheet("LongText")
    long_text.append(["record_key", "chunk_index", "chunk_count", "sha256", "chunk_text"])
    for record in long_records:
        long_text.append(
            [
                record["record_key"],
                str(record["chunk_index"]),
                str(record["chunk_count"]),
                record["sha256"],
                record["chunk_text"],
            ]
        )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.data_type = "s"
    workbook.save(path)


def _xlsx_text(
    value: str, *, record_key: str, long_records: list[dict[str, Any]]
) -> str:
    """Return inline text or a lossless descriptor backed by safe XLSX chunks."""
    utf16_units = len(value.encode("utf-16-le")) // 2
    if len(value) <= XLSX_DIRECT_TEXT_LIMIT and utf16_units <= 32_000:
        return value
    chunks = [
        value[offset : offset + XLSX_CHUNK_TEXT_LIMIT]
        for offset in range(0, len(value), XLSX_CHUNK_TEXT_LIMIT)
    ]
    digest = hashlib.sha256(value.encode("utf-8")).hexdigest()
    for index, chunk in enumerate(chunks):
        if len(chunk.encode("utf-16-le")) // 2 > 32_000:
            raise AssertionError("XLSX long-text chunk exceeds the UTF-16 safety limit")
        long_records.append(
            {
                "record_key": record_key,
                "chunk_index": index,
                "chunk_count": len(chunks),
                "sha256": digest,
                "chunk_text": chunk,
            }
        )
    return canonical_json(
        {
            "encoding": XLSX_LONG_TEXT_VERSION,
            "record_key": record_key,
            "chunk_count": len(chunks),
            "sha256": digest,
        }
    )


def _csv_safe(value: Any) -> str:
    text = value if isinstance(value, str) else canonical_json(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text
