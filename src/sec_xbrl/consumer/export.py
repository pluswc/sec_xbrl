"""Canonical snapshots and offline exports over prepared consumer responses.

This module is deliberately downstream of :class:`AnalysisClient`.  It does
not read Raw material, choose observations, calculate financial values, or
change accounting meaning.
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
from collections.abc import Iterable
from copy import deepcopy
from pathlib import Path
from typing import Any

SNAPSHOT_VERSION = "u5-consumer-export-v1"
FORMATS = {"json", "csv", "html", "xlsx"}


def prepare_overview_snapshot(
    client: Any,
    ticker: str,
    *,
    fiscal_start: int | None = None,
    fiscal_end: int | None = None,
    view: str = "LATEST_REPORTED",
    as_of: str | None = None,
    review_cutoff: str | None = None,
    recent_quarters: int = 8,
    periods: Iterable[tuple[int, int]] | None = None,
    row_ids: Iterable[str] | None = None,
) -> dict[str, Any]:
    """Return one canonical snapshot of an exact prepared overview selection."""
    response = client.overview(
        ticker,
        fiscal_start=fiscal_start,
        fiscal_end=fiscal_end,
        view=view,
        as_of=as_of,
        review_cutoff=review_cutoff,
        recent_quarters=recent_quarters,
    )
    rows = response["rows"]
    columns = response["columns"]
    cells = response["cells"]
    selected_rows = _select(rows, "row_id", row_ids, "row")
    selected_periods = _select_periods(columns, periods)
    selected_ids = {row["row_id"] for row in selected_rows}
    period_keys = {(column["fiscal_year"], column["fiscal_quarter"]) for column in selected_periods}
    cell_index = {
        (cell["row_id"], cell["fiscal_year"], cell["fiscal_quarter"], cell["period_class"]): cell
        for cell in cells
    }
    snapshot_cells = []
    for row in selected_rows:
        for column in selected_periods:
            if column.get("period_class") != row.get("period_class"):
                continue
            key = (row["row_id"], column["fiscal_year"], column["fiscal_quarter"], column["period_class"])
            source = cell_index.get(key)
            if source is None:
                source = {
                    "row_id": row["row_id"],
                    "column_id": column.get("column_id") or column.get("fiscal_time_series_column_id"),
                    "fiscal_year": column["fiscal_year"],
                    "fiscal_quarter": column["fiscal_quarter"],
                    "period_class": column["period_class"],
                    "status": "UNAVAILABLE",
                    "value": None,
                    "reason": "NO_PREPARED_CELL_FOR_SELECTED_ROW_AND_PERIOD",
                }
            snapshot_cells.append(_cell(client, response["context"], source))
    selected_metrics = [
        deepcopy(metric) for metric in response.get("metrics", [])
        if metric.get("row_id") in selected_ids
        and (metric.get("fiscal_year"), metric.get("fiscal_quarter")) in period_keys
    ]
    lineage_cells = _metric_input_cells(client, ticker, response, selected_metrics, view)
    excluded_cell_warnings = _excluded_cell_warnings(cells, selected_ids, period_keys)
    return _finish(
        kind="OVERVIEW",
        context=response["context"],
        publication={"analysis_publication_id": response["context"]["publication_id"]},
        columns=selected_periods,
        rows=selected_rows,
        cells=snapshot_cells,
        warnings=response.get("warnings", []),
        excluded_rows=[row for row in rows if row["row_id"] not in selected_ids],
        selected_periods=period_keys,
        source_identity={"profile": "core", "metrics": selected_metrics,
                         "metric_input_cells": lineage_cells},
        excluded_cell_warnings=excluded_cell_warnings,
    )


def prepare_axis_snapshot(
    client: Any,
    ticker: str,
    lens_id: str,
    *,
    context: dict[str, Any],
    periods: Iterable[tuple[int, int]] | None = None,
    row_ids: Iterable[str] | None = None,
) -> dict[str, Any]:
    """Return one canonical snapshot of an exact prepared Axis table selection."""
    response = client.axis_timeseries(ticker, lens_id, context=context)
    rows = response["rows"]
    selected_rows = _select(rows, "row_id", row_ids, "row")
    selected_columns = _select_periods(response["columns"], periods)
    period_keys = {(c["fiscal_year"], c["fiscal_quarter"]) for c in selected_columns}
    cells = []
    for row in selected_rows:
        by_period = {(c["fiscal_year"], c["fiscal_quarter"]): c for c in row["cells"]}
        for column in selected_columns:
            cells.append(_cell(client, context, by_period[(column["fiscal_year"], column["fiscal_quarter"])]))
    selected_ids = {row["row_id"] for row in selected_rows}
    return _finish(
        kind="AXIS_TIMESERIES",
        context=response["axis_context"],
        publication=response["publication"],
        columns=selected_columns,
        rows=[{key: deepcopy(value) for key, value in row.items() if key != "cells"} for row in selected_rows],
        cells=cells,
        warnings=response.get("warnings", []),
        excluded_rows=[row for row in rows if row["row_id"] not in selected_ids],
        selected_periods=period_keys,
        source_identity={
            "lens_id": lens_id,
            "lens": deepcopy(response["lens"]),
            "axis_identity": deepcopy(response["axis_identity"]),
            "reviewed_relationships": deepcopy(response["reviewed_relationships"]),
            "source_comparisons": deepcopy(response["source_comparisons"]),
        },
        excluded_cell_warnings=_excluded_cell_warnings(
            [cell for row in rows for cell in row["cells"]], selected_ids, period_keys
        ),
    )


def export_snapshot(snapshot: dict[str, Any], destination: Path | str, *, format: str | None = None) -> Path:
    """Write JSON, CSV, HTML, or XLSX from one validated canonical snapshot."""
    destination = Path(destination)
    selected_format = (format or destination.suffix.removeprefix(".")).lower()
    if selected_format not in FORMATS:
        raise ValueError(f"unsupported export format: {selected_format}")
    _validate_snapshot(snapshot)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if selected_format == "json":
        destination.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    elif selected_format == "csv":
        destination.write_text(_csv(snapshot), encoding="utf-8-sig", newline="")
    elif selected_format == "html":
        destination.write_text(_html(snapshot), encoding="utf-8")
    else:
        _xlsx(snapshot, destination)
    return destination


def _cell(client: Any, context: dict[str, Any], source: dict[str, Any]) -> dict[str, Any]:
    cell = deepcopy(source)
    provenance = cell.pop("provenance", None)
    cell_id = cell.get("cell_id")
    if provenance is None and cell_id and cell.get("status") in {"REPORTED", "DERIVED"}:
        provenance = client.trace(cell_id, context=context)["trace"]
    display_value = cell.get("value")
    if cell.get("quality_status") == "BLOCK":
        display_value = None
    raw_value = cell.pop("raw_value", None)
    return {
        **cell,
        "display_value": display_value,
        "value": display_value,
        "raw_value_lineage_only": raw_value,
        "provenance": deepcopy(provenance),
    }


def _select(rows: list[dict[str, Any]], key: str, requested: Iterable[str] | None, label: str) -> list[dict[str, Any]]:
    if requested is None:
        return deepcopy(rows)
    ids = list(requested)
    if len(ids) != len(set(ids)):
        raise ValueError(f"duplicate selected {label} identity")
    indexed = {row[key]: row for row in rows}
    missing = [identity for identity in ids if identity not in indexed]
    if missing:
        raise ValueError(f"selected {label} is not prepared: {missing}")
    return [deepcopy(indexed[identity]) for identity in ids]


def _select_periods(columns: list[dict[str, Any]], requested: Iterable[tuple[int, int]] | None) -> list[dict[str, Any]]:
    if requested is None:
        return deepcopy(columns)
    periods = [tuple(period) for period in requested]
    if len(periods) != len(set(periods)) or any(len(period) != 2 for period in periods):
        raise ValueError("selected periods must be unique fiscal year/quarter pairs")
    grouped: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for column in columns:
        grouped.setdefault((column["fiscal_year"], column["fiscal_quarter"]), []).append(column)
    missing = [period for period in periods if period not in grouped]
    if missing:
        raise ValueError(f"selected period is not prepared: {missing}")
    return [deepcopy(column) for period in periods for column in grouped[period]]


def _finish(
    *, kind: str, context: dict[str, Any], publication: dict[str, Any], columns: list[dict[str, Any]],
    rows: list[dict[str, Any]], cells: list[dict[str, Any]], warnings: list[Any],
    excluded_rows: list[dict[str, Any]], selected_periods: set[tuple[int, int]], source_identity: dict[str, Any],
    excluded_cell_warnings: list[dict[str, Any]],
) -> dict[str, Any]:
    protected = []
    for row in excluded_rows:
        row_warnings = deepcopy(row.get("warnings") or [])
        if row.get("protected") or row.get("required") or row_warnings:
            protected.append({"row_id": row["row_id"], "protected": bool(row.get("protected")),
                              "required": bool(row.get("required")), "warnings": row_warnings,
                              "reason": "EXCLUDED_BY_EXPLICIT_ROW_SELECTION"})
    snapshot = {
        "snapshot_version": SNAPSHOT_VERSION,
        "kind": kind,
        "context": deepcopy(context),
        "publication": deepcopy(publication),
        "source_identity": deepcopy(source_identity),
        "display_policy": {"unit_divisor": "1", "value_field": "display_value",
                           "raw_value_field": "raw_value_lineage_only"},
        "selection": {
            "row_ids": [row["row_id"] for row in rows],
            "periods": [list(period) for period in dict.fromkeys(
                (column["fiscal_year"], column["fiscal_quarter"]) for column in columns
            )],
            "partial_rows": bool(excluded_rows),
            "partial_periods": selected_periods != {tuple(p) for p in context["periods"]},
            "excluded_protected_or_warning_rows": protected,
            "excluded_cell_warnings": excluded_cell_warnings,
        },
        "columns": deepcopy(columns),
        "rows": deepcopy(rows),
        "cells": deepcopy(cells),
        "warnings": deepcopy(warnings),
    }
    snapshot["snapshot_id"] = _snapshot_hash(snapshot)
    return snapshot


def _excluded_cell_warnings(cells: list[dict[str, Any]], selected_rows: set[str],
                            selected_periods: set[tuple[int, int]]) -> list[dict[str, Any]]:
    result = []
    for cell in cells:
        reasons = deepcopy(cell.get("quality_reasons") or [])
        selected = cell.get("row_id") in selected_rows and (
            cell.get("fiscal_year"), cell.get("fiscal_quarter")
        ) in selected_periods
        if reasons and not selected:
            result.append({"cell_id": cell.get("cell_id"), "row_id": cell.get("row_id"),
                           "period": [cell.get("fiscal_year"), cell.get("fiscal_quarter")],
                           "warnings": reasons, "reason": "EXCLUDED_BY_EXPLICIT_SELECTION"})
    return result


def _metric_input_cells(client: Any, ticker: str, response: dict[str, Any],
                        metrics: list[dict[str, Any]], view: str) -> list[dict[str, Any]]:
    required = {identity for metric in metrics for identity in metric.get("input_ids", [])}
    if not required:
        return []
    broad = response
    company = getattr(client, "manifest", {}).get("companies", {}).get(ticker.upper(), {})
    years = company.get("years")
    if years:
        broad = client.overview(ticker, fiscal_start=min(years), fiscal_end=max(years), view=view)
    indexed = {cell.get("cell_id"): cell for cell in broad["cells"]}
    missing = sorted(required - indexed.keys())
    if missing:
        raise ValueError(f"prepared metric input cells are outside available snapshot context: {missing}")
    result = []
    for identity in sorted(required):
        cell = _cell(client, broad["context"], indexed[identity])
        cell["inputs_only"] = True
        result.append(cell)
    return result


def _snapshot_hash(snapshot: dict[str, Any]) -> str:
    material = {key: value for key, value in snapshot.items() if key != "snapshot_id"}
    canonical = json.dumps(material, sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str)
    return hashlib.sha256(canonical.encode()).hexdigest()


def _validate_snapshot(snapshot: dict[str, Any]) -> None:
    if snapshot.get("snapshot_version") != SNAPSHOT_VERSION or snapshot.get("snapshot_id") != _snapshot_hash(snapshot):
        raise ValueError("not a canonical U5 export snapshot")


def _safe(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    result = str(value)
    if result.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + result
    return result


def _flat_rows(snapshot: dict[str, Any]) -> list[list[str]]:
    row_index = {row["row_id"]: row for row in snapshot["rows"]}
    result = []
    for cell in snapshot["cells"]:
        row = row_index[cell["row_id"]]
        result.append([
            row["row_id"], row.get("label"), row.get("depth"), cell.get("fiscal_year"),
            cell.get("fiscal_quarter"), cell.get("period_class"), cell.get("display_value"),
            cell.get("status"), cell.get("quality_status"), cell.get("reason"),
            cell.get("unit"), cell.get("dimensions"), cell.get("basis_version"),
            cell.get("source_fact_id"), cell.get("cell_id"),
        ])
    return [[_safe(value) for value in row] for row in result]


HEADERS = ["row_id", "label", "depth", "fiscal_year", "fiscal_quarter", "period_class", "display_value",
           "status", "quality_status", "unavailable_reason", "unit", "dimensions", "basis_version",
           "source_fact_id", "cell_id"]


def _csv(snapshot: dict[str, Any]) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(["record_type", *HEADERS, "record_json"])
    metadata = {key: snapshot[key] for key in snapshot if key != "cells"}
    writer.writerow(["SNAPSHOT", *([""] * len(HEADERS)), _safe(metadata)])
    for row, source in zip(_flat_rows(snapshot), snapshot["cells"], strict=True):
        writer.writerow(["CELL", *row, _safe(source)])
    return output.getvalue()


def _html(snapshot: dict[str, Any]) -> str:
    esc = lambda value: html.escape(_safe(value), quote=True)
    metadata = esc(json.dumps({key: snapshot[key] for key in ("snapshot_id", "kind", "context", "publication", "selection", "warnings")}, ensure_ascii=False, indent=2))
    headings = "".join(f"<th>{esc(value)}</th>" for value in HEADERS)
    rows = "".join("<tr>" + "".join(f"<td>{esc(value)}</td>" for value in row) + "</tr>" for row in _flat_rows(snapshot))
    canonical = json.dumps(snapshot, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")
    return ("<!doctype html><html><meta charset=\"utf-8\"><title>SEC XBRL export</title>"
            "<style>body{font:14px system-ui;margin:24px}table{border-collapse:collapse}th,td{border:1px solid #bbb;padding:6px;vertical-align:top}th{background:#eee;position:sticky;top:0}pre{white-space:pre-wrap}</style>"
            f"<body><h1>SEC XBRL prepared export</h1><pre>{metadata}</pre><table><thead><tr>{headings}</tr></thead><tbody>{rows}</tbody></table>"
            f"<script type=\"application/json\" id=\"sec-xbrl-snapshot\">{canonical}</script></body></html>")


def _xlsx(snapshot: dict[str, Any], destination: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    table = workbook.active
    table.title = "Table"
    periods = list(dict.fromkeys(
        (column["fiscal_year"], column["fiscal_quarter"]) for column in snapshot["columns"]
    ))
    headers = ["label", "period_class", "unit", "row_warnings", "cell_statuses", "row_id", "depth", *[
        f"FY{year} Q{quarter}" for year, quarter in periods
    ]]
    table.append(headers)
    cell_index = {(cell["row_id"], cell["fiscal_year"], cell["fiscal_quarter"], cell["period_class"]): cell
                  for cell in snapshot["cells"]}
    for row in snapshot["rows"]:
        period_class = row.get("period_class") or snapshot.get("source_identity", {}).get("lens", {}).get("period_class")
        row_cells = [cell for cell in snapshot["cells"] if cell["row_id"] == row["row_id"]]
        if not period_class and row_cells:
            period_class = row_cells[0].get("period_class")
        units = list(dict.fromkeys(_display_unit(cell.get("unit")) for cell in row_cells))
        statuses = [{"period": [cell["fiscal_year"], cell["fiscal_quarter"]],
                     "status": cell.get("status"), "quality_status": cell.get("quality_status"),
                     "reason": cell.get("reason"), "quality_reasons": cell.get("quality_reasons") or []}
                    for cell in row_cells]
        values = [row.get("label"), period_class, " | ".join(units),
                  _json_text(row.get("warnings") or []), _json_text(statuses),
                  row["row_id"], row.get("depth")]
        for year, quarter in periods:
            cell = cell_index.get((row["row_id"], year, quarter, period_class))
            values.append(None if cell is None else cell.get("display_value"))
        table.append([_safe(value) for value in values])
    for cell in table[1]:
        cell.font = Font(bold=True)
    table.freeze_panes = "H2"
    table.auto_filter.ref = table.dimensions
    for name, width in {"A": 28, "B": 16, "C": 28, "D": 34, "E": 48, "F": 34, "G": 9}.items():
        table.column_dimensions[name].width = width
    table.column_dimensions["F"].hidden = True
    table.column_dimensions["G"].hidden = True
    for column in range(8, table.max_column + 1):
        table.column_dimensions[table.cell(1, column).column_letter].width = 20

    cells = workbook.create_sheet("Cells")
    cell_headers = sorted({key for row in snapshot["cells"] for key in row if key != "provenance"})
    cells.append(cell_headers)
    for row in snapshot["cells"]:
        cells.append([_safe(row.get(key)) for key in cell_headers])

    provenance = workbook.create_sheet("Provenance")
    provenance.append(["cell_id", "part", "parts", "provenance_json_chunk"])
    for row in snapshot["cells"]:
        _append_chunks(provenance, _safe(row.get("cell_id")), _json_text(row.get("provenance")))

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["key", "value"])
    for key in ("snapshot_version", "snapshot_id", "kind", "context", "publication", "source_identity",
                "display_policy", "selection", "warnings"):
        _append_chunks(metadata, key, _json_text(snapshot[key]))
    _append_chunks(metadata, "canonical_snapshot", _json_text(snapshot))
    for sheet in workbook.worksheets:
        if sheet.title != "Table":
            sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.data_type = "s"
    workbook.save(destination)


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _display_unit(unit: Any) -> str:
    if not isinstance(unit, list) or len(unit) != 2:
        return _json_text(unit)
    parts = []
    for value in unit:
        if isinstance(value, str):
            try:
                value = json.loads(value)
            except json.JSONDecodeError:
                value = [value] if value else []
        parts.append(value or [])
    numerator = " × ".join(map(str, parts[0])) or "1"
    denominator = " × ".join(map(str, parts[1]))
    return numerator + (" / " + denominator if denominator else "")


def _append_chunks(sheet: Any, identity: str, value: str) -> None:
    # Excel limits a cell to 32,767 characters. Keep margin for readers that
    # count UTF-16 code units and make every chunk position explicit.
    size = 16_000
    chunks = [value[index:index + size] for index in range(0, len(value), size)] or [""]
    if sheet.title == "Metadata":
        if sheet.max_column < 4:
            sheet.delete_rows(1, 1)
            sheet.append(["key", "part", "parts", "value_json_chunk"])
        for index, chunk in enumerate(chunks, 1):
            sheet.append([identity, str(index), str(len(chunks)), chunk])
    else:
        for index, chunk in enumerate(chunks, 1):
            sheet.append([identity, str(index), str(len(chunks)), chunk])
