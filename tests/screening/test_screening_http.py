from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
import threading
import urllib.error
import urllib.request
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from sec_xbrl.screening.server import create_server
from sec_xbrl.screening.service import ScreeningService
from sec_xbrl.screening.store import ScreeningStore


def _request(url: str, *, body: dict | None = None, origin: str | None = None) -> tuple[int, bytes]:
    headers = {}
    data = None
    if body is not None:
        data = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    if origin:
        headers["Origin"] = origin
    request = urllib.request.Request(url, data=data, headers=headers, method="POST" if body is not None else "GET")
    try:
        with urllib.request.urlopen(request) as response:
            return response.status, response.read()
    except urllib.error.HTTPError as exc:
        return exc.code, exc.read()


def test_loopback_http_native_ui_origin_and_restart(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    store_path = tmp_path / "journal.sqlite"
    service = ScreeningService({"fixture": screening_dataset}, ScreeningStore(store_path))
    server = create_server(service, export_root=tmp_path / "exports")
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    host = f"127.0.0.1:{server.server_port}"
    base = "http://" + host
    try:
        status, html = _request(base + "/")
        assert status == 200
        text = html.decode()
        assert "기업 스크리닝 작업대" in text and "전 회사 결과" in text
        assert "localStorage" not in text
        status, icon = _request(base + "/favicon.ico")
        assert status == 204 and icon == b""
        status, body = _request(
            base + "/api/screens",
            body={"name": "브라우저 조건", "content": revision_content},
            origin=base,
        )
        assert status == 201
        revision = json.loads(body)
        status, body = _request(
            base + "/api/runs",
            body={"revision_id": revision["revision_id"], "dataset_id": "fixture"},
            origin=base,
        )
        run = json.loads(body)
        assert status == 201 and run["counts"]["target"] == 3
        status, body = _request(
            base + f"/api/runs/{run['run_id']}/exports",
            body={"scope": "FULL", "format": "csv", "selected_ciks": None},
            origin=base,
        )
        export = json.loads(body)
        assert status == 201
        status, downloaded = _request(base + export["download_url"])
        assert status == 200 and downloaded.startswith(b"\xef\xbb\xbf")
        private = tmp_path / "exports" / "unregistered.txt"
        private.write_text("not an export", encoding="utf-8")
        status, _ = _request(base + "/exports/unregistered.txt")
        assert status == 400
        status, body = _request(
            base + "/api/screens",
            body={"name": "외부 origin", "content": revision_content},
            origin="https://example.com",
        )
        assert status == 403 and "origin" in json.loads(body)["error"]
        status, _ = _request(base + "/api/debug/fail-save")
        assert status == 404
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
    restarted = ScreeningService({"fixture": screening_dataset}, ScreeningStore(store_path))
    assert restarted.store.list_screens()[0]["revision_id"] == revision["revision_id"]


def test_server_rejects_non_loopback(screening_dataset: Path, tmp_path: Path) -> None:
    service = ScreeningService({"fixture": screening_dataset}, ScreeningStore(tmp_path / "db.sqlite"))
    with pytest.raises(ValueError, match="loopback"):
        create_server(service, export_root=tmp_path / "exports", host="0.0.0.0")
    outside = tmp_path / "outside"
    outside.mkdir()
    linked = tmp_path / "linked-exports"
    linked.symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match="symlink"):
        create_server(service, export_root=linked)


def test_http_metric_catalog_forwards_dataset_scoped_business_fields(tmp_path: Path) -> None:
    expected = [
        {
            "metric_id": "current_ratio",
            "label": "유동비율",
            "category": "FINANCIAL_STABILITY",
            "unit_kind": "PERCENT",
            "unit": "PERCENT",
            "definition_version": "screening-current-ratio-v1",
            "period_class": "INSTANT",
            "coverage": {
                "companies_available": 5,
                "companies_total": 21,
                "periods_available": 17,
            },
            "business_availability": "AVAILABLE",
            "unavailable_reasons": ["MISSING_INPUT"],
        }
    ]

    def metric_catalog(dataset_id: str) -> list[dict]:
        assert dataset_id == "expanded-v2"
        return expected

    service = SimpleNamespace(metric_catalog=metric_catalog)
    server = create_server(service, export_root=tmp_path / "exports")
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        status, body = _request(
            f"http://127.0.0.1:{server.server_port}/api/metric-catalog?dataset_id=expanded-v2"
        )
        assert status == 200
        assert json.loads(body) == expected
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_http_repeated_and_selected_exports_remain_immutable(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    store_path = tmp_path / "journal.sqlite"
    export_root = tmp_path / "exports"
    service = ScreeningService({"fixture": screening_dataset}, ScreeningStore(store_path))
    revision = service.create_screen("반복 내보내기", revision_content)
    run = service.execute(revision["revision_id"], "fixture")
    expected_rows = {row["cik"]: row for row in run["rows"]}
    server = create_server(service, export_root=export_root)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_port}"
    exports: list[tuple[dict, bytes]] = []
    try:
        for selected_format in ("csv", "xlsx"):
            for _repeat in range(2):
                exports.append(
                    _http_export(
                        base,
                        run["run_id"],
                        scope="FULL",
                        selected_format=selected_format,
                        selected_ciks=None,
                    )
                )
        for selected_format, selected_ciks in (
            ("csv", ["0000000001"]),
            ("xlsx", ["0000000002"]),
        ):
            exports.append(
                _http_export(
                    base,
                    run["run_id"],
                    scope="SELECTED",
                    selected_format=selected_format,
                    selected_ciks=selected_ciks,
                )
            )

        urls = [manifest["download_url"] for manifest, _data in exports]
        assert len(urls) == len(set(urls)) == 6
        assert all(
            re.fullmatch(
                rf"/exports/{re.escape(run['run_id'])}-(?:full|selected)-"
                r"(?:csv|xlsx)-[0-9a-f]{32}\.(?:csv|xlsx)",
                url,
            )
            for url in urls
        )
        for manifest, data in exports:
            scope, rows = _export_payload(manifest["format"], data)
            expected_ciks = (
                set(expected_rows)
                if manifest["scope"] == "FULL"
                else set(manifest["selected_ciks"])
            )
            assert scope == manifest["scope"]
            assert {row["cik"] for row in rows} == expected_ciks
            assert rows == [row for row in run["rows"] if row["cik"] in expected_ciks]

        for manifest, original in exports:
            status, downloaded = _request(base + manifest["download_url"])
            assert status == 200 and downloaded == original
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)

    with sqlite3.connect(store_path) as connection:
        stored = [
            json.loads(row[0])
            for row in connection.execute(
                "SELECT manifest_json FROM export_manifests ORDER BY rowid"
            )
        ]
    assert len(stored) == 6
    assert len({manifest["path"] for manifest in stored}) == 6
    assert all(Path(manifest["path"]).is_file() for manifest in stored)
    assert [manifest["scope"] for manifest in stored].count("FULL") == 4
    assert [manifest["scope"] for manifest in stored].count("SELECTED") == 2


def test_save_failure_has_no_false_commit(
    screening_dataset: Path,
    revision_content: dict,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = ScreeningService({"fixture": screening_dataset}, ScreeningStore(tmp_path / "db.sqlite"))

    class BrokenConnection:
        def __enter__(self):
            raise OSError("synthetic commit failure")

        def __exit__(self, *args):
            return False

    monkeypatch.setattr(service.store, "_connect", lambda: BrokenConnection())
    with pytest.raises(OSError, match="synthetic"):
        service.create_screen("저장 실패", revision_content)
    recovered = ScreeningStore(tmp_path / "db.sqlite")
    assert recovered.list_screens() == []


def _http_export(
    base: str,
    run_id: str,
    *,
    scope: str,
    selected_format: str,
    selected_ciks: list[str] | None,
) -> tuple[dict, bytes]:
    status, body = _request(
        base + f"/api/runs/{run_id}/exports",
        body={
            "scope": scope,
            "format": selected_format,
            "selected_ciks": selected_ciks,
        },
        origin=base,
    )
    assert status == 201
    manifest = json.loads(body)
    status, downloaded = _request(base + manifest["download_url"])
    assert status == 200
    return manifest, downloaded


def _export_payload(selected_format: str, data: bytes) -> tuple[str, list[dict]]:
    if selected_format == "csv":
        records = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        header = json.loads(records[0]["record_json"])
        rows = [
            json.loads(record["record_json"])
            for record in records
            if record["record_type"] == "ROW"
        ]
        return header["metadata"]["scope"], rows
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    metadata = {
        row[0].value: row[1].value
        for row in workbook["Metadata"].iter_rows(min_row=2)
    }
    headers = [cell.value for cell in workbook["Rows"][1]]
    record_index = headers.index("record_json")
    rows = [
        json.loads(row[record_index].value)
        for row in workbook["Rows"].iter_rows(min_row=2)
    ]
    return metadata["scope"], rows
