from __future__ import annotations

import copy
from pathlib import Path

import pytest

from sec_xbrl.screening.common import digest_value
from sec_xbrl.screening.service import _validate_period_policy
from sec_xbrl.screening.store import _revision_content

ANNUAL_REVISION = {
    "universe": {"mode": "ALL_PREPARED"},
    "predicates": [
        {
            "metric_id": "revenue",
            "operator": "GTE",
            "threshold": "50000000000",
            "definition_version": "annual-core-reported-v1",
        }
    ],
    "period_policy": {"mode": "FIXED_FISCAL_YEAR", "fiscal_year": 2025},
    "columns": ["ticker", "name", "sic_code", "aggregate", "predicates"],
    "sort": {"column": "ticker", "direction": "ASC"},
}


def test_quarter_policy_is_additive_and_annual_revision_identity_is_unchanged() -> None:
    expected = "9048ff2a509e2d19cd3d9cef5f23ed94473fcde9936bc70626d6971f99ab8e82"
    original = copy.deepcopy(ANNUAL_REVISION)

    _validate_period_policy(original["period_policy"])
    normalized = _revision_content(original)

    assert normalized == ANNUAL_REVISION
    assert digest_value(normalized) == expected
    _validate_period_policy(
        {
            "mode": "FIXED_CALENDAR_QUARTER_END",
            "calendar_year": 2026,
            "calendar_quarter": 2,
        }
    )


@pytest.mark.parametrize(
    "policy",
    [
        {"mode": "FIXED_CALENDAR_QUARTER_END", "calendar_year": 2026},
        {"mode": "FIXED_CALENDAR_QUARTER_END", "calendar_year": 2026, "calendar_quarter": 0},
        {"mode": "FIXED_CALENDAR_QUARTER_END", "calendar_year": True, "calendar_quarter": 2},
        {
            "mode": "FIXED_CALENDAR_QUARTER_END",
            "calendar_year": 2026,
            "calendar_quarter": 2,
            "fiscal_quarter": 2,
        },
    ],
)
def test_quarter_policy_rejects_incomplete_or_fiscal_label_inputs(
    policy: dict[str, object],
) -> None:
    with pytest.raises(ValueError, match="calendar quarter-end"):
        _validate_period_policy(policy)


from sec_xbrl.screening.quarterly import QuarterlyScreeningMetricReader


def test_wire_v2_is_rejected_for_cross_record_source_contradictions() -> None:
    root = Path(__file__).parents[1] / "fixtures" / "screening_quarterly_wire_v2"
    with pytest.raises(ValueError, match="source|filing|view"):
        QuarterlyScreeningMetricReader(
            root,
            manifest_sha256="fa743d5d24a228e7f4fd5c541127361f17046f3574184c4b2f17de72252143e3",
        )


def test_superseded_wire_fixture_rejects_noncanonical_ttm_namespace() -> None:
    with pytest.raises(ValueError, match="prefixed"):
        QuarterlyScreeningMetricReader(
            "/tmp/universe50-quarterly-schema-v1",
            manifest_sha256="ff57954804c9866f568903580e368021e0a3ad27f1e3e7ce76fddafa476825f5",
        )


def test_authoritative_quarter_wire_v5_is_verified() -> None:
    root = Path(__file__).parents[1] / "fixtures" / "screening_quarterly_wire_v5"
    reader = QuarterlyScreeningMetricReader(
        root,
        manifest_sha256="84e652322e570f4768c278fa924aaf6c1de98f1158be0d162b9be0a488fa359a",
    )
    status = reader.verify_all()
    assert status["observations"] == 64
    assert status["period_inventory"] == 16
    assert {row["metric_id"] for row in reader.definitions} == {
        "quarter_revenue",
        "quarter_cfo",
        "quarter_gross_profit",
        "ttm_revenue",
    }
    march_end = next(
        row
        for row in reader.observations
        if row["cik"] == "0000000001"
        and row["metric_id"] == "quarter_revenue"
        and (row["calendar_year"], row["calendar_quarter"]) == (2025, 1)
    )
    assert march_end["period_end"] == "2025-04-01"


import csv
import json

from openpyxl import load_workbook

from sec_xbrl.screening.dataset_v3 import ScreeningDatasetV3Reader
from sec_xbrl.screening.service import ScreeningService
from sec_xbrl.screening.store import ScreeningStore


def _quarter_service(tmp_path: Path) -> ScreeningService:
    wire = QuarterlyScreeningMetricReader(
        Path(__file__).parents[1] / "fixtures" / "screening_quarterly_wire_v5"
    )
    reader = object.__new__(ScreeningDatasetV3Reader)
    reader.publication_id = "synthetic-v3-dataset"
    reader.manifest = {
        "dataset_id": "synthetic-v3",
        "fixture_kind": "SYNTHETIC_QUARTER_SCHEMA_FIXTURE",
        "financial_as_of": "2026-10-10",
        "sources": {"metadata": {"publication_id": "synthetic-metadata"}},
        "versions": {"quarterly_publication_id": wire.publication_id},
    }
    reader.companies = [
        {"cik": cik, "ticker": f"SYNTH{i}", "name": f"Synthetic {i}", "sic_code": "0000"}
        for i, cik in enumerate(wire.manifest["universe"]["ciks"], 1)
    ]
    reader.metric_definitions = []
    reader.metric_support = []
    reader.observations = []
    reader.quarterly = wire
    reader.quarter_definitions = wire.definitions
    reader.quarter_support = wire.support
    reader.quarter_observations = wire.observations
    reader.quarter_period_inventory = wire.period_inventory
    reader.company = lambda cik: next(
        copy.deepcopy(row) for row in reader.companies if row["cik"] == cik
    )
    service = object.__new__(ScreeningService)
    service.datasets = {"synthetic-v3": reader}
    service.explorers = {}
    service.company_descriptions = None
    service.store = ScreeningStore(tmp_path / "quarter-journal.sqlite")
    return service


def test_quarter_run_trace_comparison_and_export_preserve_period_identity(tmp_path: Path) -> None:
    service = _quarter_service(tmp_path)
    content = {
        "universe": {"mode": "ALL_PREPARED"},
        "predicates": [
            {
                "metric_id": "quarter_revenue",
                "operator": "GTE",
                "threshold": "1",
                "definition_version": "quarter-reported-core-v1",
            }
        ],
        "period_policy": {
            "mode": "FIXED_CALENDAR_QUARTER_END",
            "calendar_year": 2025,
            "calendar_quarter": 1,
        },
        "columns": ["company", "aggregate", "predicates"],
        "sort": {"column": "ticker", "direction": "ASC"},
    }
    revision = service.create_screen("분기", content)
    run = service.execute(revision["revision_id"], "synthetic-v3")
    assert run["period_policy"] == content["period_policy"]
    first = run["rows"][0]["predicates"][0]["observation"]
    assert first["period_end"] == "2025-04-01"
    assert first["calendar_quarter"] == 1
    trace = service.trace(run["run_id"], cik=run["rows"][0]["cik"], metric_id="quarter_revenue")
    assert trace["trace"] == first["source_lineage"]
    comparison = service.quarter_comparison(
        "synthetic-v3",
        ciks=[row["cik"] for row in run["rows"]],
        metric_id="quarter_revenue",
        calendar_year=2025,
        calendar_quarter=1,
    )
    assert len(comparison["rows"]) == 2
    exported = tmp_path / "quarter.json"
    service.export_run(run["run_id"], exported)
    payload = json.loads(exported.read_text())
    exported_observation = payload["rows"][0]["predicates"][0]["observation"]
    assert exported_observation["fiscal_quarter"] == first["fiscal_quarter"]
    assert exported_observation["selection_as_of_date"] == "2026-10-10"
    csv_path = tmp_path / "quarter.csv"
    service.export_run(run["run_id"], csv_path, format="csv")
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
    csv_observation = json.loads(csv_rows[1]["record_json"])["predicates"][0]["observation"]
    assert csv_observation["calendar_quarter"] == 1
    assert csv_observation["period_end"] == "2025-04-01"
    xlsx_path = tmp_path / "quarter.xlsx"
    service.export_run(run["run_id"], xlsx_path, format="xlsx")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    record = workbook["Rows"].cell(row=2, column=7).value
    assert '"calendar_quarter":1' in record
    assert '"selection_as_of_date":"2026-10-10"' in record


def test_quarter_explorer_keeps_all_companies_and_explicit_gaps(tmp_path: Path) -> None:
    service = _quarter_service(tmp_path)
    result = service.quarter_explorer("synthetic-v3", calendar_year=2026, calendar_quarter=2)
    assert len(result["rows"]) == 2
    assert len(result["metrics"]) == 4
    assert all(row["period"]["calendar_quarter"] == 2 for row in result["rows"])
    held = result["rows"][0]["metrics"]["ttm_revenue"]
    assert held["value"] is None
    assert held["reason"] == "INPUTS_NOT_COMPATIBLE"


def test_quality_block_and_selected_bucket_coverage_are_preserved(tmp_path: Path) -> None:
    service = _quarter_service(tmp_path)
    row = next(
        x
        for x in service.datasets["synthetic-v3"].quarter_observations
        if x["metric_id"] == "quarter_revenue"
        and x["cik"] == "0000000001"
        and (x["calendar_year"], x["calendar_quarter"]) == (2025, 1)
    )
    row.update(quality_status="BLOCK", quality_blocking=True)
    content = copy.deepcopy(ANNUAL_REVISION)
    content["predicates"][0].update(
        metric_id="quarter_revenue", threshold="1", definition_version="quarter-reported-core-v1"
    )
    content["period_policy"] = {
        "mode": "FIXED_CALENDAR_QUARTER_END",
        "calendar_year": 2025,
        "calendar_quarter": 1,
    }
    revision = service.create_screen("quality", content)
    result = service.execute(revision["revision_id"], "synthetic-v3")["rows"][0]
    assert result["aggregate"] == "UNDETERMINED"
    assert result["quality_blocking"] is True
    explorer = service.quarter_explorer("synthetic-v3", calendar_year=2024, calendar_quarter=4)
    cfo = next(x for x in explorer["metrics"] if x["metric_id"] == "quarter_cfo")
    assert cfo["coverage"]["companies_available"] == 0
    assert cfo["aggregate_coverage"]["periods_available"] == 2


def test_v3_binding_requires_exact_annual_metric_parent_and_cutoff() -> None:
    import hashlib
    from types import SimpleNamespace

    from sec_xbrl.screening.dataset_v3 import _verify_parent_binding

    ciks = ["0000000001", "0000000002"]
    membership = hashlib.sha256(("\n".join(ciks) + "\n").encode()).hexdigest()
    annual = SimpleNamespace(
        companies=[{"cik": x} for x in ciks],
        manifest={
            "financial_as_of": "2026-10-10",
            "sources": {
                "metrics": {"publication_id": "annual-metrics", "manifest_sha256": "a" * 64}
            },
        },
    )
    quarter = SimpleNamespace(
        manifest={
            "financial_as_of": "2026-10-10",
            "annual_parent": {
                "schema_version": "screening-metric-publication-v1",
                "publication_id": "annual-metrics",
                "manifest_sha256": "a" * 64,
                "company_membership_sha256": membership,
            },
        }
    )
    _verify_parent_binding(annual, quarter)
    quarter.manifest["annual_parent"]["publication_id"] = "unrelated"
    with pytest.raises(ValueError, match="annual metric parent"):
        _verify_parent_binding(annual, quarter)


@pytest.mark.parametrize("case", ["future_filing", "support_count", "membership"])
def test_rehashed_semantic_corruption_is_rejected(tmp_path: Path, case: str) -> None:
    import hashlib
    import json
    import shutil

    source = Path(__file__).parents[1] / "fixtures" / "screening_quarterly_wire_v5"
    root = tmp_path / case
    shutil.copytree(source, root)
    manifest_path = root / "quarterly_screening_metric_manifest.json"
    manifest = json.loads(manifest_path.read_text())
    if case == "membership":
        manifest["universe"]["company_membership_sha256"] = "f" * 64
        manifest["annual_parent"]["company_membership_sha256"] = "f" * 64
    else:
        filename = (
            "quarter_metric_observations.json"
            if case == "future_filing"
            else "quarter_metric_support.json"
        )
        payload = json.loads((root / filename).read_text())
        if case == "future_filing":
            payload["records"][0]["filed_date"] = "2099-01-01"
        else:
            payload["records"][0]["available_bucket_count"] = 999
        data = (json.dumps(payload, indent=2) + "\n").encode()
        (root / filename).write_bytes(data)
        manifest["files"][filename].update(sha256=hashlib.sha256(data).hexdigest(), size=len(data))
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
    with pytest.raises(ValueError):
        QuarterlyScreeningMetricReader(root)


def test_v4_wire_is_preserved_but_rejects_implicit_available_scope() -> None:
    root = Path(__file__).parents[1] / "fixtures" / "screening_quarterly_wire_v4"
    with pytest.raises(ValueError, match="view|scope"):
        QuarterlyScreeningMetricReader(
            root, manifest_sha256="aad6f74c6fdfe23f0663b4adc5b12b49d8b22af09ea130ec06a316e589bcd9c2"
        )
