from __future__ import annotations

import csv
import hashlib
import json
import shutil
import sqlite3
from collections.abc import Callable
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from sec_xbrl.screening import dataset as dataset_module
from sec_xbrl.screening.common import digest_value, sha256_file, write_json
from sec_xbrl.screening.dataset import (
    COMPANIES,
    MANIFEST,
    OBSERVATIONS,
    ScreeningDatasetReader,
    _is_usd_unit,
    _observation,
    _verify_analysis_files,
)
from sec_xbrl.screening.service import (
    XLSX_DIRECT_TEXT_LIMIT,
    XLSX_LONG_TEXT_VERSION,
    ScreeningService,
    _write_xlsx,
    _xlsx_text,
)
from sec_xbrl.screening.store import ScreeningStore, backup_store, restore_store


def _service(dataset: Path, tmp_path: Path) -> ScreeningService:
    return ScreeningService({"fixture": dataset}, ScreeningStore(tmp_path / "journal.sqlite"))


def test_executor_preserves_every_target_unknown_and_original_quality(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    revision = service.create_screen("FY2025 예시", revision_content)
    run = service.execute(revision["revision_id"], "fixture")
    assert run["counts"] == {"target": 3, "pass": 1, "fail": 1, "undetermined": 1, "quality_block": 0}
    assert [(row["ticker"], row["aggregate"]) for row in run["rows"]] == [
        ("ONE", "PASS"), ("THREE", "UNDETERMINED"), ("TWO", "FAIL")
    ]
    unknown = next(row for row in run["rows"] if row["ticker"] == "THREE")
    assert unknown["predicates"][0]["reason"] == "MULTIPLE_CANDIDATES"
    assert unknown["predicates"][0]["observation"]["value"] is None
    warning = next(row for row in run["rows"] if row["ticker"] == "TWO")
    assert warning["predicates"][0]["observation"]["quality_status"] == "WARN"
    assert warning["quality_blocking"] is False
    assert service.store.run(run["run_id"])["output_hash"] == run["output_hash"]


def test_unknown_and_unsupported_metrics_fail_closed(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    for metric in ("not_a_metric", "revenue_growth"):
        invalid = json.loads(json.dumps(revision_content))
        invalid["predicates"] = [{"metric_id": metric, "operator": "GTE", "threshold": "0",
                                  "definition_version": "annual-core-reported-v1"}]
        revision = service.create_screen(metric, invalid)
        with pytest.raises(ValueError, match="unknown|unavailable"):
            service.execute(revision["revision_id"], "fixture")
    invalid_number = json.loads(json.dumps(revision_content))
    invalid_number["predicates"][0]["threshold"] = "NaN"
    with pytest.raises(ValueError, match="finite"):
        service.create_screen("bad", invalid_number)
    wrong_version = json.loads(json.dumps(revision_content))
    wrong_version["predicates"][0]["definition_version"] = "invented-definition-v99"
    revision = service.create_screen("wrong-version", wrong_version)
    with pytest.raises(ValueError, match="definition version"):
        service.execute(revision["revision_id"], "fixture")


def test_sic_universe_metadata_unknown_trace_and_peer_dates(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    narrowed = json.loads(json.dumps(revision_content))
    narrowed["universe"] = {"mode": "SIC", "sic_codes": ["3674"]}
    revision = service.create_screen("반도체", narrowed)
    run = service.execute(revision["revision_id"], "fixture")
    assert [row["cik"] for row in run["rows"]] == ["0000000001", "0000000002"]
    inventory = service.metadata_inventory("fixture", prepared=False)
    assert inventory["total"] == 1 and inventory["rows"][0]["sic_status"] == "UNKNOWN_SIC"
    peers = service.peer_candidates("fixture", cik="0000000001")
    assert [row["cik"] for row in peers["peers"]] == ["0000000002"]
    trace = service.trace(run["run_id"], cik="0000000001", metric_id="revenue")
    assert trace["trace"]["value_lineage"]["source_filing_id"] == "0000000001-25-000001"
    comparison = service.compare_peers(
        "fixture", ciks=["0000000001", "0000000002"], metric_id="revenue", fiscal_year=2025
    )
    assert all(row["observation"]["context_end_convention"] == "SOURCE_CONTEXT_END_EXCLUSIVE" for row in comparison["rows"])


def test_append_only_decisions_restart_backup_and_blank_restore(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    revision = service.create_screen("영속 조건", revision_content)
    run = service.execute(revision["revision_id"], "fixture")
    first = service.add_decision(cik="0000000001", state="INTEREST", rationale="원천 확인",
                                 run_id=run["run_id"], actor="tester", evidence_ids=["fact-1"])
    second = service.add_decision(cik="0000000001", state="HOLD", rationale="추가 검토",
                                  run_id=run["run_id"], actor="tester", evidence_ids=["fact-2"])
    assert second["previous_event_id"] == first["event_id"]
    assert second["before_state"] == "INTEREST"
    restarted = _service(screening_dataset, tmp_path)
    assert restarted.store.revision(revision["revision_id"])["content"] == revision_content
    assert [event["state"] for event in restarted.decisions("0000000001")] == ["INTEREST", "HOLD"]
    backup = backup_store(tmp_path / "journal.sqlite", tmp_path / "backup")
    restored_path = restore_store(backup, tmp_path / "restored.sqlite")
    restored = ScreeningStore(restored_path)
    assert restored.run(run["run_id"])["output_hash"] == run["output_hash"]
    assert len(restored.decisions("0000000001")) == 2
    with pytest.raises(ValueError, match="new destination"):
        restore_store(backup, restored_path)


def test_run_comparison_separates_revision_from_universe_membership(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    first_revision = service.create_screen("비교", revision_content)
    first = service.execute(first_revision["revision_id"], "fixture")
    changed = json.loads(json.dumps(revision_content))
    changed["columns"] = [*changed["columns"], "decision"]
    second_revision = service.revise_screen(first_revision["screen_id"], changed)
    second = service.execute(second_revision["revision_id"], "fixture")
    comparison = service.compare_runs(first["run_id"], second["run_id"])
    assert comparison["changed_dimensions"]["condition_revision"] is True
    assert comparison["changed_dimensions"]["universe"] is False
    assert comparison["changed_dimensions"]["universe_selector"] is False
    assert comparison["result_state"] == "UNCHANGED"
    assert comparison["universe_membership"]["old_ciks"] == comparison[
        "universe_membership"
    ]["new_ciks"]
    assert second["versions"]["predicate_definition_versions"] == [
        "annual-core-reported-v1"
    ]


def test_run_read_reconciles_counts_and_universe_rows(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    revision = service.create_screen("무결성", revision_content)
    run = service.execute(revision["revision_id"], "fixture")
    altered = json.loads(json.dumps(run))
    altered.pop("run_id")
    altered.pop("created_at")
    altered["counts"]["target"] = 999
    with sqlite3.connect(service.store.path) as connection:
        connection.execute(
            "UPDATE screen_runs SET payload_json=? WHERE run_id=?",
            (json.dumps(altered), run["run_id"]),
        )
    with pytest.raises(ValueError, match="counts do not reconcile"):
        service.store.run(run["run_id"])


def test_full_and_selected_csv_xlsx_equal_frozen_run(
    screening_dataset: Path, revision_content: dict, tmp_path: Path
) -> None:
    service = _service(screening_dataset, tmp_path)
    revision = service.create_screen("내보내기", revision_content)
    run = service.execute(revision["revision_id"], "fixture")
    outputs = {}
    for scope, selected in (("FULL", None), ("SELECTED", ["0000000001"])):
        for extension in ("csv", "xlsx"):
            destination = tmp_path / f"{scope.lower()}.{extension}"
            outputs[(scope, extension)] = service.export_run(
                run["run_id"], destination, scope=scope, selected_ciks=selected
            )
    full_rows = list(csv.DictReader((tmp_path / "full.csv").open(encoding="utf-8-sig")))
    assert [json.loads(row["record_json"])["cik"] for row in full_rows if row["record_type"] == "ROW"] == [row["cik"] for row in run["rows"]]
    selected_rows = list(csv.DictReader((tmp_path / "selected.csv").open(encoding="utf-8-sig")))
    assert [row["cik"] for row in selected_rows if row["record_type"] == "ROW"] == ["0000000001"]
    for scope, records in (("FULL", full_rows), ("SELECTED", selected_rows)):
        header = json.loads(records[0]["record_json"])
        rebuilt = {
            **header["metadata"],
            "rows": [
                json.loads(row["record_json"])
                for row in records
                if row["record_type"] == "ROW"
            ],
        }
        assert digest_value(rebuilt) == header["canonical_payload_hash"]
        assert digest_value(rebuilt) == outputs[(scope, "csv")]["canonical_payload_hash"]
    for scope, expected in (("full", 3), ("selected", 1)):
        workbook = load_workbook(tmp_path / f"{scope}.xlsx", data_only=False)
        assert workbook.sheetnames == ["Rows", "Metadata", "LongText"]
        assert workbook["Rows"].max_row == expected + 1
        assert all(cell.data_type != "f" for sheet in workbook for row in sheet.iter_rows() for cell in row)
        resolve = _xlsx_resolver(workbook)
        metadata = {
            row[0].value: resolve(row[1].value)
            for row in workbook["Metadata"].iter_rows(min_row=2)
        }
        headers = [cell.value for cell in workbook["Rows"][1]]
        records = [
            dict(zip(headers, (cell.value for cell in row), strict=True))
            for row in workbook["Rows"].iter_rows(min_row=2)
        ]
        rebuilt = {
            "export_version": metadata["export_version"],
            "scope": metadata["scope"],
            "selected_ciks": json.loads(metadata["selected_ciks_json"]),
            "run": json.loads(metadata["run_json"]),
            "rows": [json.loads(resolve(row["record_json"])) for row in records],
        }
        assert metadata["xlsx_long_text_encoding"] == XLSX_LONG_TEXT_VERSION
        assert digest_value(rebuilt) == metadata["canonical_payload_hash"]
        assert digest_value(rebuilt) == outputs[(scope.upper(), "xlsx")][
            "canonical_payload_hash"
        ]
        assert outputs[(scope.upper(), "xlsx")]["xlsx_encoding_version"] == XLSX_LONG_TEXT_VERSION
    assert all(result["canonical_payload_hash"] for result in outputs.values())
    full_path = tmp_path / "full.csv"
    assert service.store.export_for_path(full_path)["row_count"] == 3
    unregistered = tmp_path / "not-an-export.csv"
    unregistered.write_text("private file", encoding="utf-8")
    with pytest.raises(ValueError, match="not uniquely registered"):
        service.store.export_for_path(unregistered)
    full_path.write_text("changed", encoding="utf-8")
    with pytest.raises(ValueError, match="integrity mismatch"):
        service.store.export_for_path(full_path)


def test_xlsx_long_text_chunks_reassemble_full_and_selected_payloads(tmp_path: Path) -> None:
    long_text = ("장기근거😀" * 7_000) + "끝"
    rows = [
        {
            "cik": f"{index + 1:010d}",
            "ticker": f"T{index}",
            "name": f"기업 {index}",
            "sic_code": "3571",
            "aggregate": "PASS",
            "quality_blocking": False,
            "predicates": [{"metric_id": "revenue", "trace": long_text + str(index)}],
        }
        for index in range(21)
    ]
    base = {
        "export_version": "screen-run-export-v1",
        "selected_ciks": None,
        "run": {"run_id": "run-long", "provenance": long_text},
    }
    for suffix, selected_rows in (("full", rows), ("selected", rows[:2]), ("repeat", rows)):
        payload = {
            **base,
            "scope": "FULL" if suffix != "selected" else "SELECTED",
            "selected_ciks": None if suffix != "selected" else [row["cik"] for row in selected_rows],
            "rows": selected_rows,
        }
        destination = tmp_path / f"{suffix}.xlsx"
        _write_xlsx(destination, payload)
        workbook = load_workbook(destination, data_only=False)
        resolve = _xlsx_resolver(workbook)
        metadata = {
            row[0].value: resolve(row[1].value)
            for row in workbook["Metadata"].iter_rows(min_row=2)
        }
        headers = [cell.value for cell in workbook["Rows"][1]]
        records = [
            dict(zip(headers, (cell.value for cell in row), strict=True))
            for row in workbook["Rows"].iter_rows(min_row=2)
        ]
        rebuilt = {
            "export_version": metadata["export_version"],
            "scope": metadata["scope"],
            "selected_ciks": json.loads(metadata["selected_ciks_json"]),
            "run": json.loads(metadata["run_json"]),
            "rows": [json.loads(resolve(row["record_json"])) for row in records],
        }
        assert rebuilt == payload
        assert metadata["canonical_payload_hash"] == digest_value(payload)
        assert workbook["LongText"].max_row > 1
        assert all(
            len(str(cell.value or "").encode("utf-16-le")) // 2 <= 32_000
            for sheet in workbook
            for row in sheet.iter_rows()
            for cell in row
        )


def test_xlsx_keeps_fitting_legacy_text_inline() -> None:
    text = "a" * XLSX_DIRECT_TEXT_LIMIT
    long_records: list[dict] = []

    assert _xlsx_text(text, record_key="row:0:0000000001", long_records=long_records) == text
    assert long_records == []


def _xlsx_resolver(workbook) -> Callable[[str], str]:
    grouped: dict[str, list[tuple[int, int, str, str]]] = {}
    for row in workbook["LongText"].iter_rows(min_row=2, values_only=True):
        record_key, chunk_index, chunk_count, digest, chunk_text = row
        grouped.setdefault(record_key, []).append(
            (int(chunk_index), int(chunk_count), digest, chunk_text)
        )

    def resolve(value: str) -> str:
        try:
            descriptor = json.loads(value)
        except (json.JSONDecodeError, TypeError):
            return value
        if not isinstance(descriptor, dict) or descriptor.get("encoding") != XLSX_LONG_TEXT_VERSION:
            return value
        chunks = sorted(grouped[descriptor["record_key"]])
        assert [row[0] for row in chunks] == list(range(descriptor["chunk_count"]))
        assert all(row[1] == descriptor["chunk_count"] for row in chunks)
        assert all(row[2] == descriptor["sha256"] for row in chunks)
        rebuilt = "".join(row[3] for row in chunks)
        assert hashlib.sha256(rebuilt.encode("utf-8")).hexdigest() == descriptor["sha256"]
        return rebuilt

    return resolve


def test_export_registration_failure_leaves_no_downloadable_file(
    screening_dataset: Path,
    revision_content: dict,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = _service(screening_dataset, tmp_path)
    revision = service.create_screen("내보내기 실패", revision_content)
    run = service.execute(revision["revision_id"], "fixture")
    destination = tmp_path / "failed.csv"

    def fail_registration(manifest: dict) -> dict:
        raise OSError("synthetic export manifest failure")

    monkeypatch.setattr(service.store, "record_export", fail_registration)
    with pytest.raises(OSError, match="synthetic"):
        service.export_run(run["run_id"], destination)
    assert not destination.exists()
    assert not list(tmp_path.glob(".failed.csv.partial-*"))


def test_dataset_reader_rejects_rehashed_record_count_or_payload_tamper(
    screening_dataset: Path,
) -> None:
    reader = ScreeningDatasetReader(screening_dataset)
    assert reader.verify_all()["counts"]["metadata_candidates"] == 4
    path = screening_dataset / "metric_observations.json"
    path.write_text("[]\n", encoding="utf-8")
    with pytest.raises(ValueError, match="integrity"):
        ScreeningDatasetReader(screening_dataset)


def test_open_reader_never_returns_later_corrupted_snapshot_bytes(
    screening_dataset: Path,
) -> None:
    reader = ScreeningDatasetReader(screening_dataset)
    cik = "0000000001"
    verified = reader.source_snapshot(cik)
    path = screening_dataset / "source_snapshots" / cik / (
        verified["snapshot_id"] + ".json"
    )
    changed = json.loads(path.read_text())
    changed["cells"][0]["display_value"] = "999999"
    write_json(path, changed)
    assert reader.source_snapshot(cik) == verified
    with pytest.raises(ValueError, match="changed after reader verification"):
        reader.verify_all()


def test_actual_dataset_cannot_omit_governed_source_anchors(screening_dataset: Path) -> None:
    path = screening_dataset / MANIFEST
    manifest = json.loads(path.read_text())
    manifest["fixture_kind"] = "ACTUAL_PREPARED"
    manifest["fixture_policy"] = None
    identity = {key: value for key, value in manifest.items() if key != "publication_id"}
    manifest["publication_id"] = "screening-dataset-" + digest_value(identity)[:24]
    write_json(path, manifest)
    with pytest.raises(ValueError, match="source anchors"):
        ScreeningDatasetReader(screening_dataset)


def test_analysis_parent_bytes_must_match_declared_hash(tmp_path: Path) -> None:
    root = tmp_path / "analysis"
    root.mkdir()
    data = root / "core.parquet"
    data.write_bytes(b"governed prepared bytes")
    manifest = {
        "companies": {
            "ONE": {
                "views": {
                    "LATEST_REPORTED": {
                        "files": {
                            "core": {"path": "core.parquet", "sha256": sha256_file(data)}
                        }
                    }
                }
            }
        }
    }
    _verify_analysis_files(root, manifest)
    data.write_bytes(b"changed while manifest stayed unchanged")
    with pytest.raises(ValueError, match="analysis data integrity mismatch"):
        _verify_analysis_files(root, manifest)


def test_usd_screening_unit_requires_exact_prepared_measure() -> None:
    assert _is_usd_unit([["iso4217:USD"], []]) is True
    assert _is_usd_unit(['["iso4217:USD"]', "[]"]) is True
    assert _is_usd_unit([["iso4217:USDollar"], []]) is False
    assert _is_usd_unit(['["iso4217:USDollar"]', "[]"]) is False
    assert _is_usd_unit([["iso4217:USD"], ["xbrli:shares"]]) is False


def test_coordinated_embedded_rehash_cannot_replace_prepared_source_value(
    screening_dataset: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Self-consistent copied JSON still has to equal the independent prepared parent."""
    publication = tmp_path / "actual"
    shutil.copytree(screening_dataset, publication)
    companies = json.loads((publication / COMPANIES).read_text())
    original_snapshots = {}
    for company in companies:
        path = publication / "source_snapshots" / company["cik"] / (
            company["source_snapshot_id"] + ".json"
        )
        original_snapshots[company["ticker"]] = json.loads(path.read_text())

    config_path = tmp_path / "config.json"
    completion_root = tmp_path / "completion"
    metadata_root = tmp_path / "metadata"
    completion_root.mkdir()
    metadata_root.mkdir()
    write_json(completion_root / "universe_delivery_manifest.json", {"fixture": True})
    write_json(metadata_root / "company_metadata_manifest.json", {"fixture": True})
    analysis_sources = []
    config_companies = []
    fake_clients = {}
    for company in companies:
        analysis_root = tmp_path / f"analysis-{company['ticker']}"
        history_root = tmp_path / f"history-{company['ticker']}"
        analysis_root.mkdir()
        history_root.mkdir()
        history = history_root / "history_manifest.json"
        write_json(history, {"ticker": company["ticker"]})
        analysis_manifest = {
            "publication_id": company["analysis_publication_id"],
            "period_mode": "ANNUAL",
            "companies": {
                company["ticker"]: {
                    "as_of": "2026-10-08",
                    "source_publication": str(history_root),
                    "views": {},
                }
            },
            "source_manifest_hashes": {"history_manifest_sha256": sha256_file(history)},
        }
        write_json(analysis_root / "analysis_manifest.json", analysis_manifest)
        fake_clients[str(analysis_root.resolve())] = SimpleNamespace(
            manifest=analysis_manifest,
            overview_snapshot=lambda *args, _snapshot=original_snapshots[company["ticker"]], **kwargs: _snapshot,
        )
        config_companies.append(
            {"ticker": company["ticker"], "cik": company["cik"], "analysis_root": str(analysis_root)}
        )
        analysis_sources.append(
            {
                "ticker": company["ticker"],
                "cik": company["cik"],
                "root": str(analysis_root.resolve()),
                "manifest_sha256": sha256_file(analysis_root / "analysis_manifest.json"),
                "publication_id": company["analysis_publication_id"],
                "source_publication": str(history_root),
                "history_manifest_sha256": sha256_file(history),
            }
        )
    config = {
        "schema_version": "screening-prepare-config-v1",
        "dataset_id": "fixture",
        "fixture_kind": "ACTUAL_PREPARED",
        "financial_as_of": "2026-10-08",
        "completion_root": str(completion_root),
        "metadata_root": str(metadata_root),
        "companies": config_companies,
    }
    write_json(config_path, config)
    metadata_rows = json.loads((publication / "metadata_inventory.json").read_text())
    candidates = [
        SimpleNamespace(
            cik=row["cik"],
            name=row["name"],
            tickers=tuple(row["tickers"]),
            exchanges=tuple(row["exchanges"]),
            sic_code=row["sic_code"],
            sic_description=row["sic_description"],
            metadata_status=row["metadata_status"],
            observed_at=row["observed_at"],
        )
        for row in metadata_rows
    ]
    completion_rows = [
        {
            "cik": company["cik"],
            "ticker": company["ticker"],
            "collection_state": "READY",
            "held_filings": 1,
            "selected_filings": 1,
            "pending_filings": 0,
            "failed_filings": 0,
            "industry_status": "CLASSIFIED",
        }
        for company in companies
    ]
    monkeypatch.setattr(
        dataset_module,
        "UniverseDeliveryReader",
        lambda root: SimpleNamespace(
            publication_id="completion-fixture",
            manifest={
                "sources": {
                    "accuracy_assessment": {
                        "state": "PASS",
                        "parser_version": "fixture-parser-v1",
                        "parser_rule_version": "fixture-parser-rule-v1",
                        "frozen_commit_sha": "fixture-frozen-commit",
                        "sha256": "fixture-assessment-hash",
                    }
                }
            },
            readiness={
                "assessment_state": "READY",
                "dates": {"financial_as_of": "2026-10-08"},
                "companies": completion_rows,
            },
        ),
    )
    monkeypatch.setattr(
        dataset_module,
        "CompanyMetadataReader",
        lambda root: SimpleNamespace(
            publication_id="metadata-fixture", list_companies=lambda: candidates
        ),
    )
    monkeypatch.setattr(
        dataset_module, "open_analysis", lambda root: fake_clients[str(root.resolve())]
    )

    manifest_path = publication / MANIFEST
    manifest = json.loads(manifest_path.read_text())
    manifest["fixture_kind"] = "ACTUAL_PREPARED"
    manifest["fixture_policy"] = None
    manifest["sources"] = {
        "config": {"path": str(config_path), "sha256": sha256_file(config_path)},
        "completion": {
            "root": str(completion_root),
            "publication_id": "completion-fixture",
            "manifest_sha256": sha256_file(
                completion_root / "universe_delivery_manifest.json"
            ),
        },
        "metadata": {
            "root": str(metadata_root),
            "publication_id": "metadata-fixture",
            "manifest_sha256": sha256_file(metadata_root / "company_metadata_manifest.json"),
        },
        "analysis": analysis_sources,
    }
    version_sets = {
        "analysis_versions": set(),
        "source_publications": {row["source_publication"] for row in analysis_sources},
        "review_cutoffs": set(),
        "source_observed_times": {row["observed_at"] for row in metadata_rows},
        "metric_definition_versions": {
            row["definition_version"]
            for row in json.loads((publication / "metric_support.json").read_text())
        },
        "concept_mapping_versions": {"fixture-mapping-v1"},
        "selection_rule_versions": {"fixture-selection-v1"},
        "classification_rule_versions": {"fixture-period-v1"},
        "accession_version_ledger_ids": set(),
        "raw_concept_taxonomy_versions": set(),
    }
    manifest["versions"] = dataset_module._version_manifest(
        version_sets,
        {
            "parser_version": "fixture-parser-v1",
            "parser_rule_version": "fixture-parser-rule-v1",
            "frozen_commit_sha": "fixture-frozen-commit",
            "sha256": "fixture-assessment-hash",
        },
    )

    target = companies[0]
    old_path = publication / "source_snapshots" / target["cik"] / (
        target["source_snapshot_id"] + ".json"
    )
    changed_snapshot = json.loads(old_path.read_text())
    changed_snapshot["cells"][0]["display_value"] = "999"
    changed_snapshot["cells"][0]["value"] = "999"
    material = {key: value for key, value in changed_snapshot.items() if key != "snapshot_id"}
    changed_snapshot["snapshot_id"] = digest_value(material)
    new_path = old_path.with_name(changed_snapshot["snapshot_id"] + ".json")
    old_path.unlink()
    write_json(new_path, changed_snapshot)
    target["source_snapshot_id"] = changed_snapshot["snapshot_id"]
    write_json(publication / COMPANIES, companies)
    observations = json.loads((publication / OBSERVATIONS).read_text())
    observations = [row for row in observations if row["cik"] != target["cik"]]
    observations.extend(
        _observation(target["cik"], target["ticker"], cell["row_id"], cell, changed_snapshot)
        for cell in changed_snapshot["cells"]
    )
    write_json(publication / OBSERVATIONS, observations)
    declared_files = {
        str(path.relative_to(publication)): {
            "sha256": sha256_file(path),
            "size": path.stat().st_size,
        }
        for path in publication.rglob("*")
        if path.is_file() and path != manifest_path
    }
    manifest["files"] = declared_files
    identity = {key: value for key, value in manifest.items() if key != "publication_id"}
    manifest["publication_id"] = "screening-dataset-" + digest_value(identity)[:24]
    write_json(manifest_path, manifest)

    with pytest.raises(ValueError, match="differs from prepared analysis source"):
        ScreeningDatasetReader(publication)
