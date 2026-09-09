from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import ClassVar

from openpyxl import load_workbook

from sec_xbrl.consumer.export import (
    export_snapshot,
    prepare_axis_snapshot,
    prepare_overview_snapshot,
)


class PreparedClient:
    context: ClassVar[dict] = {
        "ticker": "TEST", "view": "LATEST_REPORTED", "as_of": "2026-09-10",
        "review_cutoff": "2026-09-10T00:00:00Z", "periods": [[2026, 1], [2026, 2]],
        "publication_id": "analysis-publication", "source_publication": "history-publication",
        "context_id": "prepared-context",
    }

    def overview(self, ticker, **kwargs):
        assert ticker == "TEST"
        return {
            "context": self.context,
            "rows": [
                {"row_id": "safe", "label": "Revenue", "period_class": "QTD_3M"},
                {"row_id": "warning", "label": "=warning", "period_class": "QTD_3M", "protected": True,
                 "warnings": ["MANDATORY_WARNING"]},
            ],
            "columns": [
                {"column_id": "q1", "fiscal_year": 2026, "fiscal_quarter": 1, "period_class": "QTD_3M"},
                {"column_id": "q2", "fiscal_year": 2026, "fiscal_quarter": 2, "period_class": "QTD_3M"},
            ],
            "cells": [
                {"cell_id": "reported", "row_id": "safe", "fiscal_year": 2026, "fiscal_quarter": 1,
                 "period_class": "QTD_3M", "status": "REPORTED", "value": "123456789012345678.90",
                 "raw_value": "123456789012345678.90", "unit": [["iso4217:USD"], []], "source_fact_id": "fact-1"},
                {"cell_id": "blocked", "row_id": "safe", "fiscal_year": 2026, "fiscal_quarter": 2,
                 "period_class": "QTD_3M", "status": "REPORTED", "value": None,
                 "raw_value": "12559938000", "quality_status": "BLOCK", "quality_reasons": ["REVIEW_BLOCK"]},
                {"cell_id": "formula", "row_id": "warning", "fiscal_year": 2026, "fiscal_quarter": 1,
                 "period_class": "QTD_3M", "status": "DERIVED", "value": "=1+1"},
                {"cell_id": "negative", "row_id": "warning", "fiscal_year": 2026, "fiscal_quarter": 2,
                 "period_class": "QTD_3M", "status": "REPORTED", "value": "-263653000"},
            ],
            "metrics": [{"metric_id": "QOQ", "row_id": "safe", "fiscal_year": 2026,
                         "fiscal_quarter": 2, "formula": "(current/prior)-1",
                         "input_ids": ["reported", "blocked"], "status": "UNAVAILABLE"}],
            "warnings": ["GLOBAL_WARNING"],
        }

    def trace(self, cell_id, *, context):
        assert context == self.context
        return {"trace": {"cell_id": cell_id, "value_lineage": {"formula": "=SUM(A1:A2)",
                        "source_inputs": [{"fact_id": "input", "unit_denominator_measures": '["shares"]'}]}}}

    def axis_timeseries(self, ticker, lens_id, *, context):
        base = self.overview(ticker)
        cells = []
        for row_id in ("total", "member"):
            cells.append({**base["cells"][0], "row_id": row_id, "cell_id": row_id + "-q1",
                          "provenance": self.trace(row_id + "-q1", context=context)["trace"]})
            cells.append({**base["cells"][1], "row_id": row_id, "cell_id": row_id + "-q2",
                          "provenance": self.trace(row_id + "-q2", context=context)["trace"]})
        return {
            "axis_context": {**context, "lens_id": lens_id}, "context": context,
            "publication": {"analysis_publication_id": "analysis-publication", "axis_publication_id": "axis-publication"},
            "lens": {"node_id": lens_id, "label": "Products"}, "axis_identity": [["axis", "member"]],
            "period_class": "QTD_3M", "columns": base["columns"],
            "rows": [
                {"row_id": "total", "label": "Total", "depth": 0, "protected": True, "warnings": [], "cells": cells[:2]},
                {"row_id": "member", "label": "Member", "depth": 1, "warnings": ["ROW_WARNING"], "cells": cells[2:]},
            ],
            "reviewed_relationships": [{"formula": "=member", "input_ids": ["member-q1"]}],
            "source_comparisons": [{"review_id": "review"}], "warnings": ["AXIS_WARNING"],
        }


def test_overview_partial_snapshot_preserves_order_null_and_provenance():
    snapshot = prepare_overview_snapshot(
        PreparedClient(), "TEST", periods=[(2026, 2), (2026, 1)], row_ids=["safe"]
    )
    assert snapshot["selection"]["row_ids"] == ["safe"]
    assert snapshot["selection"]["periods"] == [[2026, 2], [2026, 1]]
    assert [cell["cell_id"] for cell in snapshot["cells"]] == ["blocked", "reported"]
    blocked = snapshot["cells"][0]
    assert blocked["display_value"] is None and blocked["raw_value_lineage_only"] == "12559938000"
    assert snapshot["cells"][1]["display_value"] == "123456789012345678.90"
    assert snapshot["cells"][1]["provenance"]["value_lineage"]["source_inputs"][0]["unit_denominator_measures"] == '["shares"]'
    assert snapshot["warnings"] == ["GLOBAL_WARNING"]
    assert snapshot["source_identity"]["metrics"][0]["input_ids"] == ["reported", "blocked"]
    assert [cell["cell_id"] for cell in snapshot["source_identity"]["metric_input_cells"]] == ["blocked", "reported"]
    assert snapshot["selection"]["excluded_protected_or_warning_rows"] == [{
        "row_id": "warning", "protected": True, "required": False, "warnings": ["MANDATORY_WARNING"],
        "reason": "EXCLUDED_BY_EXPLICIT_ROW_SELECTION",
    }]


def test_all_formats_share_snapshot_and_are_spreadsheet_injection_safe(tmp_path: Path):
    snapshot = prepare_overview_snapshot(PreparedClient(), "TEST")
    snapshot["source_identity"]["large_evidence"] = "한" * 32_100 + "@boundary"
    snapshot["snapshot_id"] = __import__("hashlib").sha256(json.dumps(
        {key: value for key, value in snapshot.items() if key != "snapshot_id"},
        sort_keys=True, ensure_ascii=False, separators=(",", ":"), default=str,
    ).encode()).hexdigest()
    paths = {extension: export_snapshot(snapshot, tmp_path / f"export.{extension}")
             for extension in ("json", "csv", "html", "xlsx")}
    assert json.loads(paths["json"].read_text())["snapshot_id"] == snapshot["snapshot_id"]
    rows = list(csv.DictReader(io.StringIO(paths["csv"].read_text(encoding="utf-8-sig"))))
    assert next(row for row in rows if row["cell_id"] == "reported")["display_value"] == "123456789012345678.90"
    assert next(row for row in rows if row["cell_id"] == "blocked")["display_value"] == ""
    assert "<td>=1+1</td>" in paths["html"].read_text()
    assert "<td>-263653000</td>" in paths["html"].read_text()
    embedded = paths["html"].read_text().split('id="sec-xbrl-snapshot">', 1)[1].split("</script>", 1)[0]
    assert json.loads(embedded) == snapshot
    workbook = load_workbook(paths["xlsx"], data_only=False)
    assert workbook.sheetnames == ["Table", "Cells", "Provenance", "Metadata"]
    values = [[cell.value for cell in row] for row in workbook["Table"].iter_rows()]
    assert any("123456789012345678.90" in row for row in values)
    assert any("=1+1" in row for row in values)
    assert any("-263653000" in row for row in values)
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet.iter_rows() for cell in row)
    assert "source_inputs" in workbook["Provenance"][2][3].value
    metadata_rows = list(workbook["Metadata"].iter_rows(min_row=2, values_only=True))
    canonical = "".join(row[3] for row in metadata_rows if row[0] == "canonical_snapshot")
    assert json.loads(canonical) == snapshot
    csv_metadata = next(row for row in rows if row["record_type"] == "SNAPSHOT")
    assert json.loads(csv_metadata["record_json"])["rows"] == snapshot["rows"]


def test_axis_snapshot_retains_publication_review_and_warning_evidence():
    client = PreparedClient()
    snapshot = prepare_axis_snapshot(client, "TEST", "lens", context=client.context,
                                     periods=[(2026, 1)], row_ids=["member"])
    assert snapshot["publication"]["axis_publication_id"] == "axis-publication"
    assert snapshot["source_identity"]["reviewed_relationships"][0]["input_ids"] == ["member-q1"]
    assert snapshot["source_identity"]["source_comparisons"] == [{"review_id": "review"}]
    assert snapshot["warnings"] == ["AXIS_WARNING"]
    assert snapshot["selection"]["excluded_protected_or_warning_rows"][0]["row_id"] == "total"


def test_export_rejects_mutation_duplicate_selection_and_missing_trace(tmp_path: Path):
    client = PreparedClient()
    snapshot = prepare_overview_snapshot(client, "TEST")
    snapshot["warnings"].append("MUTATED")
    try:
        export_snapshot(snapshot, tmp_path / "bad.json")
        raise AssertionError("mutated hash should fail")
    except ValueError as exc:
        assert "canonical" in str(exc)
    try:
        prepare_overview_snapshot(client, "TEST", row_ids=["safe", "safe"])
        raise AssertionError("duplicate rows should fail")
    except ValueError as exc:
        assert "duplicate" in str(exc)

    class MissingTrace(PreparedClient):
        def trace(self, cell_id, *, context):
            raise ValueError("missing")

    try:
        prepare_overview_snapshot(MissingTrace(), "TEST")
        raise AssertionError("reported cell without trace should fail closed")
    except ValueError as exc:
        assert "missing" in str(exc)

    class NullTrace(PreparedClient):
        def trace(self, cell_id, *, context):
            return {"trace": None}

    class MismatchedTrace(PreparedClient):
        def trace(self, cell_id, *, context):
            return {"trace": {"cell_id": "other", "value_lineage": {}}}

    for broken in (NullTrace(), MismatchedTrace()):
        try:
            prepare_overview_snapshot(broken, "TEST")
            raise AssertionError("null or mismatched trace should fail closed")
        except ValueError as exc:
            assert "trace" in str(exc)

    class MissingCellId(PreparedClient):
        def overview(self, ticker, **kwargs):
            response = super().overview(ticker, **kwargs)
            response["cells"][0].pop("cell_id")
            return response

    try:
        prepare_overview_snapshot(MissingCellId(), "TEST")
        raise AssertionError("reported cell without identity should fail closed")
    except ValueError as exc:
        assert "cell_id" in str(exc)
